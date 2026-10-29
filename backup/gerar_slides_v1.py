#!/usr/bin/env python3
"""gerar_slides.py — monta slides.pptx (10 slides, ~8 min) a partir dos CSVs e PNGs.

Os números vêm de dados/ (tabela_esperas.csv, resumo.csv, testes); as notas do
apresentador dividem a fala entre [Integrante 1], [Integrante 2] e [Integrante 3].
Uso: python3 scripts/gerar_slides.py   (requer python-pptx e pandas)
"""
import textwrap
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
D, G = RAIZ / "dados", RAIZ / "graficos"

VERDE = RGBColor(0x1F, 0x4D, 0x3A)      # dominante
VERDE_CLARO = RGBColor(0xE6, 0xEF, 0xEA)
VERMELHO = RGBColor(0xB3, 0x1B, 0x2B)   # acento
TINTA = RGBColor(0x1A, 0x1A, 0x1A)
CINZA = RGBColor(0x55, 0x55, 0x55)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE_T, FONTE_C, FONTE_M = "Cambria", "Calibri", "Courier New"


def br(x, n=2):
    return f"{x:.{n}f}".replace(".", ",")


# ------------------------------------------------------------------ dados
res = pd.read_csv(D / "resumo.csv")
tab = pd.read_csv(D / "tabela_esperas.csv").set_index("configuracao").to_dict("index")
n_exec, n_ok = len(res), int((res["status"] == "ok").sum())
viol = int(res["violacoes_prioridade"].sum() + res["violacoes_regras"].sum())
est = pd.read_csv(D / "estresse" / "resumo.csv")
ing = pd.read_csv(D / "ingenua" / "resumo.csv").iloc[0]
ing_ev = pd.read_csv(D / "ingenua" / "eventos.csv")
tp = pd.read_csv(D / "teste_prioridade" / "resumo.csv").iloc[0]
saida_teste = (D / "teste_prioridade" / "saida.txt").read_text().strip().splitlines()
r03, r06, r12, a05, d30 = (tab[k] for k in ("E03_R10", "E06_R10", "E12_R10", "E06_R05", "E06_R30"))


def t_ev(df, ev, g=None):
    q = df[df["evento"] == ev]
    if g is not None:
        q = q[q["grupo_id"] == g]
    return float(q["t_ms"].iloc[0])


ing_dec = ing_ev[ing_ev["evento"] == "decisao"]["t_ms"].tolist()

# ------------------------------------------------------------------ helpers
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BRANCO_LAYOUT = prs.slide_layouts[6]


def novo(fundo=BRANCO):
    s = prs.slides.add_slide(BRANCO_LAYOUT)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = fundo
    return s


def texto(s, x, y, w, h, linhas, tam=16, cor=TINTA, fonte=FONTE_C, negrito=False, alinh=PP_ALIGN.LEFT,
          ancora=MSO_ANCHOR.TOP, nome=None, espaco=4):
    """linhas: str ou lista de str / (str, dict) com opções por parágrafo."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nome:
        tb.name = nome
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(linhas, str):
        linhas = [linhas]
    for i, l in enumerate(linhas):
        op = {}
        if isinstance(l, tuple):
            l, op = l
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = op.get("alinh", alinh)
        p.space_after = Pt(op.get("espaco", espaco))
        r = p.add_run()
        r.text = ("•  " if op.get("bala") else "") + l
        f = r.font
        f.size = Pt(op.get("tam", tam))
        f.bold = op.get("negrito", negrito)
        f.italic = op.get("italico", False)
        f.name = op.get("fonte", fonte)
        f.color.rgb = op.get("cor", cor)
    return tb


def titulo(s, txt, cor=VERDE):
    texto(s, 0.6, 0.35, 12.1, 0.9, txt, tam=34, cor=cor, fonte=FONTE_T, negrito=True, nome="Título")


def caixa(s, x, y, w, h, cor=VERDE_CLARO, borda=None, forma=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = cor
    if borda is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borda
    sh.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.08
    return sh


def circulo(s, x, y, d, txt, cor=VERDE, tam=20):
    c = caixa(s, x, y, d, d, cor=cor, forma=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(tam), True, BRANCO, FONTE_C
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def imagem(s, arq, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(arq), Inches(x), Inches(y), **kw)


def notas(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def rodape(s, n, escuro=False):
    texto(s, 0.6, 7.0, 9, 0.35, "Tema 11 — Problema do Papai Noel · Sistemas Operacionais (FAG)", tam=10,
          cor=BRANCO if escuro else CINZA)
    texto(s, 11.9, 7.0, 0.8, 0.35, str(n), tam=10, cor=BRANCO if escuro else CINZA, alinh=PP_ALIGN.RIGHT)


# ------------------------------------------------------------------ 1. capa
s = novo(VERDE)
texto(s, 0.9, 1.7, 11.5, 1.6, "Problema do Papai Noel", tam=54, cor=BRANCO, fonte=FONTE_T, negrito=True)
texto(s, 0.9, 2.9, 11.5, 0.9, "Sincronização com pthreads: implementação, experimentos e análise", tam=24,
      cor=BRANCO)
texto(s, 0.9, 4.5, 11.5, 1.4, ["Tema 11 · Sistemas Operacionais · Prof. Lucca Abbado Neres · FAG",
                               "[Integrante 1] · [Integrante 2] · [Integrante 3]"], tam=18, cor=BRANCO, espaco=8)
for i, (n, rot) in enumerate([("1", "Noel"), ("9", "renas"), ("N", "elfos")]):
    circulo(s, 9.4 + i * 1.15, 5.75, 0.95, n, cor=VERMELHO, tam=24)
    texto(s, 9.3 + i * 1.15, 6.72, 1.15, 0.35, rot, tam=12, cor=BRANCO, alinh=PP_ALIGN.CENTER)
notas(s, "[Integrante 1] (~15 s) Apresentar o grupo e o tema: o Problema do Papai Noel, um clássico de "
         "sincronização. Implementamos em C com pthreads, rodamos 25 execuções e vamos mostrar código, "
         "demonstração e resultados.")

# ------------------------------------------------------------------ 2. problema
s = novo()
titulo(s, "O problema: duas demandas, um único Noel")
caixa(s, 0.6, 1.5, 5.9, 4.4)
circulo(s, 0.95, 1.85, 0.9, "3", tam=26)
texto(s, 2.05, 1.9, 4.3, 0.9, "Grupos de exatamente 3 elfos", tam=20, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
texto(s, 0.95, 3.0, 5.3, 3.3, [
    ("Cada elfo faz um pedido (versão finita)", {"bala": True}),
    ("O elfo i chega em 10 + i ms", {"bala": True}),
    ("Ajudar um grupo dura 3 ms e não pode ser interrompido", {"bala": True}),
    ("Nenhum elfo pode invadir um grupo fechado ou em atendimento", {"bala": True}),
], tam=18, espaco=10)
caixa(s, 6.85, 1.5, 5.9, 4.4)
circulo(s, 7.2, 1.85, 0.9, "9", cor=VERMELHO, tam=26)
texto(s, 8.3, 1.9, 4.3, 0.9, "Nove renas, uma viagem", tam=20, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
texto(s, 7.2, 3.0, 5.3, 3.3, [
    ("Chegam juntas: antes (5 ms), junto (10 ms) ou depois (30 ms)", {"bala": True}),
    ("A viagem só é preparada com as 9 renas (5 ms)", {"bala": True}),
    ("Prioridade: com as duas demandas prontas, Noel escolhe as renas", {"bala": True}),
    ("O programa encerra quando a viagem e todos os grupos terminam", {"bala": True}),
], tam=18, espaco=10)
rodape(s, 2)
notas(s, "[Integrante 1] (~45 s) Noel dorme até ser acordado. Elfos só são atendidos em grupos de três, "
         "e um quarto elfo não pode entrar num grupo já formado. As nove renas precisam estar todas presentes "
         "para a viagem. Se as duas demandas estiverem prontas, a regra manda escolher as renas. Na versão "
         "finita há uma viagem e um pedido por elfo, e o programa termina quando tudo foi feito. Os tempos "
         "são fixos: elfos a partir de 10 ms, ajuda de 3 ms, viagem de 5 ms.")

# ------------------------------------------------------------------ 3. solução
s = novo()
titulo(s, "Solução: mutex, cond vars e fila de grupos")
blocos = [("Noel", "1 thread", "dorme em cv_noel (timedwait, CLOCK_MONOTONIC); decide com o mutex: renas primeiro", VERMELHO),
          ("Renas", "9 threads", "a 9ª grava t_nona e sinaliza Noel; esperam cv_renas até a viagem acabar", VERDE),
          ("Elfos", "N threads", "entram no grupo aberto; o 3º fecha o grupo, dá o id e sinaliza Noel; esperam cv_elfos", VERDE)]
for i, (nome, qtd, desc, cor) in enumerate(blocos):
    x = 0.6 + i * 4.1
    caixa(s, x, 1.5, 3.8, 2.6)
    texto(s, x + 0.25, 1.65, 3.4, 0.55, nome, tam=22, negrito=True, cor=cor)
    texto(s, x + 0.25, 2.15, 3.4, 0.4, qtd, tam=14, cor=CINZA)
    texto(s, x + 0.25, 2.6, 3.4, 1.5, desc, tam=14)
texto(s, 0.6, 4.4, 12.1, 2.4, [
    ("Sinal de início comum: a main espera todas as threads ficarem prontas, grava t0 e faz broadcast em cv_largada", {"bala": True}),
    ("Sem espera ocupada: os sleeps só simulam chegada, ajuda e viagem; toda a coordenação usa mutex e cond vars", {"bala": True}),
    ("Eventos em vetores na memória; os CSVs são gravados depois dos joins, e os mecanismos são destruídos no fim", {"bala": True}),
    ("Invariantes checadas na execução: grupo ≠ 3, viagem < 9 renas, prioridade, interrupção, atendido 0 ou 2 vezes", {"bala": True}),
    ("Timeout de segurança de 2000 ms: grava status \"timeout\" em vez de travar", {"bala": True}),
], tam=16, espaco=7)
rodape(s, 3)
notas(s, "[Integrante 2] (~1 min) São 1 thread Noel, 9 renas e N elfos, todas criadas antes do primeiro join. "
         "Todo o estado é protegido por um único mutex. Noel dorme numa variável de condição com relógio "
         "monotônico e timeout de segurança. A nona rena e o terceiro elfo de cada grupo acordam Noel. "
         "O tempo só começa a contar depois do sinal de início comum. Os sleeps não sincronizam nada, só "
         "simulam as durações. Os eventos ficam em memória e só vão para CSV depois dos joins.")

# ------------------------------------------------------------------ 4. grupos e decisão
s = novo()
titulo(s, "Formação de grupos e decisão de Noel")
caixa(s, 0.6, 1.5, 6.3, 5.2, cor=RGBColor(0xF3, 0xF4, 0xF6))
texto(s, 0.85, 1.65, 5.9, 5.0, [
    ("// Elfo i (com o mutex)", {"cor": CINZA}),
    "g = grupo_aberto; inscreve em g",
    "se g.n == 3:",
    "    g.estado = PRONTO",
    "    grupo_aberto++",
    "    signal(cv_noel)",
    "espera g == CONCLUÍDO",
    ("", {}),
    ("// Noel (com o mutex)", {"cor": CINZA}),
    "renas_ok = renas == 9 e !viagem",
    "elfos_ok = há grupo PRONTO",
    "se renas_ok: preparar viagem",
    "senão se elfos_ok: ajudar o próximo",
    "senão: timedwait(cv_noel)",
], tam=14, fonte=FONTE_M, espaco=1)
texto(s, 7.3, 1.6, 5.4, 0.5, "Por que funciona", tam=22, negrito=True, cor=VERDE)
texto(s, 7.3, 2.25, 5.4, 4.4, [
    ("Inscrever-se e fechar o grupo acontecem na mesma seção crítica: o 4º elfo já cai no grupo seguinte", {"bala": True}),
    ("Grupos têm id e são atendidos em FIFO, então não há starvation entre elfos", {"bala": True}),
    ("renas_ok é testado antes de elfos_ok, dentro do lock em que Noel decide", {"bala": True}),
    ("Noel simula a ajuda fora do lock, e ninguém além dele decide: nenhuma ajuda é interrompida", {"bala": True}),
    ("Sem deadlock: há um só lock, as esperas ficam em laços com predicado e cada mudança tem o seu signal", {"bala": True}),
], tam=16, espaco=9)
rodape(s, 4)
notas(s, "[Integrante 2] (~1 min) O ponto-chave é que o terceiro elfo fecha o grupo e avança o ponteiro do "
         "grupo aberto na mesma seção crítica em que se inscreveu. Então o quarto elfo nunca enxerga o grupo "
         "fechado. Os grupos ganham um id e formam uma fila FIFO. Na decisão, Noel olha primeiro as renas. "
         "Não há deadlock: existe um único lock, toda espera reavalia o predicado e toda mudança de estado "
         "vem acompanhada de signal ou broadcast.")

# ------------------------------------------------------------------ 5. ordem problemática
s = novo()
titulo(s, "A solução ingênua falha: elfos antes das renas")
linhas = [("t (ms)", "evento"),
          (br(t_ev(ing_ev, "grupo_pronto", 1), 3), "grupo 1 pronto"),
          (br(t_ev(ing_ev, "grupo_pronto", 2), 3), "grupo 2 pronto"),
          (br(t_ev(ing_ev, "nona_rena"), 3), "9ª rena chega: as duas demandas prontas"),
          (br(ing_dec[0], 3), "Noel ajuda o grupo 1  ✗ violação"),
          (br(ing_dec[1], 3), "Noel ajuda o grupo 2  ✗ violação"),
          (br(t_ev(ing_ev, "inicio_preparacao"), 3), "só agora prepara a viagem")]
tb = s.shapes.add_table(len(linhas), 2, Inches(0.6), Inches(1.55), Inches(7.2), Inches(4.3)).table
tb.columns[0].width, tb.columns[1].width = Inches(1.6), Inches(5.6)
for i, (a, b) in enumerate(linhas):
    for j, v in enumerate((a, b)):
        c = tb.cell(i, j)
        c.text = v
        f = c.text_frame.paragraphs[0].runs[0].font
        f.size, f.name = Pt(16), FONTE_C
        f.bold = i == 0 or "violação" in b
        f.color.rgb = BRANCO if i == 0 else (VERMELHO if "violação" in b else TINTA)
        c.fill.solid()
        c.fill.fore_color.rgb = VERDE if i == 0 else (BRANCO if i % 2 else VERDE_CLARO)
texto(s, 0.6, 6.05, 7.2, 0.7, "Execução real: ./simulacao --elfos 6 --renas-ms 30 --teste-prioridade --ingenua",
      tam=12, fonte=FONTE_M, cor=CINZA)
caixa(s, 8.3, 1.55, 4.45, 2.3)
texto(s, 8.5, 1.65, 4.1, 1.1, f"{int(ing['violacoes_prioridade'])} violações", tam=40, negrito=True, cor=VERMELHO,
      fonte=FONTE_T)
texto(s, 8.5, 2.75, 4.1, 1.0, f"resposta às renas: {br(ing['resposta_renas_ms'], 2)} ms "
      f"(contra {br(tp['resposta_renas_ms'] * 1000, 0)} µs na solução correta)", tam=15)
caixa(s, 8.3, 4.1, 4.45, 2.6)
texto(s, 8.5, 4.2, 4.1, 2.4, [
    ("Invasão (exemplo construído)", {"negrito": True, "cor": VERDE}),
    ("Com um contador simples, o elfo 3 incrementa para 4 antes de Noel acordar, e Noel atende 4 elfos. "
     "Evitado porque o grupo é fechado dentro do lock.", {"tam": 14}),
], tam=16, espaco=5)
rodape(s, 5)
notas(s, "[Integrante 2] (~30 s) Rodamos de verdade uma versão ingênua, que atende elfos primeiro. Os dois "
         "grupos estavam prontos quando a nona rena chegou, e Noel atendeu os dois antes das renas: o próprio "
         "programa contou duas violações de prioridade, e as renas esperaram mais de 6 ms em vez de "
         "microssegundos. O outro erro clássico, o quarto elfo invadir o grupo, é evitado porque o grupo "
         "fecha na mesma seção crítica em que o terceiro elfo se inscreve.")

# ------------------------------------------------------------------ 6. demonstração
s = novo()
titulo(s, "Demonstração: disputa real pela decisão")
caixa(s, 0.6, 1.5, 12.15, 3.75, cor=RGBColor(0x1E, 0x1E, 0x1E))
linhas_saida = []
for l in saida_teste:
    linhas_saida += textwrap.wrap(l, 96, subsequent_indent="    ", break_on_hyphens=False) or [""]
texto(s, 0.85, 1.65, 11.7, 3.6, [(l, {"cor": RGBColor(0x8F, 0xE3, 0xA8) if "OK" in l or "PASSOU" in l
                                     else RGBColor(0xE8, 0xE8, 0xE8)}) for l in linhas_saida],
      tam=14, fonte=FONTE_M, espaco=1)
passos = [("1", "Noel fica \"atrasado\": espera em cv_noel até as duas demandas ficarem prontas"),
          ("2", "Grupos prontos em 12 e 15 ms; 9ª rena em ~30 ms; a 1ª decisão é a viagem"),
          ("3", f"{int(tp['concluidos'])}/{int(tp['previstos'])} elfos atendidos, 0 violações, saída em dados/teste_prioridade/")]
for i, (n, txt) in enumerate(passos):
    x = 0.6 + i * 4.1
    circulo(s, x, 5.65, 0.6, n, cor=VERMELHO, tam=16)
    texto(s, x + 0.75, 5.55, 3.25, 1.3, txt, tam=14, ancora=MSO_ANCHOR.MIDDLE)
rodape(s, 6)
notas(s, "[Integrante 2] (~2 min) Rodar ao vivo: cd src && ./simulacao --elfos 6 --renas-ms 30 --rep 1 "
         "--semente 11 --out ../dados/teste_prioridade/ --teste-prioridade. Neste modo Noel só é liberado "
         "quando as renas e pelo menos um grupo estão prontos, o que cria a disputa de propósito. A saída mostra "
         "disputa = SIM, primeira decisão = renas, 6 de 6 elfos e zero violações. Se der tempo, rodar também "
         "a referência (--elfos 6 --renas-ms 10) e mostrar os CSVs em dados/.")

# ------------------------------------------------------------------ 7. G1
s = novo()
titulo(s, "Mais elfos, mais vazão: o custo fixo é diluído")
imagem(s, G / "G1_vazao_elfos.png", 0.6, 1.4, h=5.4)
texto(s, 9.0, 1.6, 3.8, 1.1, f"{br(r12['vazao_media_pedidos_s'], 0)}", tam=54, negrito=True, cor=VERDE, fonte=FONTE_T)
texto(s, 9.0, 2.65, 3.8, 0.6, "pedidos/s com 12 elfos", tam=16, cor=CINZA)
texto(s, 9.0, 3.4, 3.8, 3.3, [
    (f"3 → 6 → 12 elfos: {br(r03['vazao_media_pedidos_s'], 0)} → {br(r06['vazao_media_pedidos_s'], 0)} → "
     f"{br(r12['vazao_media_pedidos_s'], 0)} pedidos/s", {"bala": True}),
    ("Previsão T ≈ 15 + 3·N/3 ms: hipótese confirmada", {"bala": True}),
    ("Diferenças > 100 pedidos/s ≫ faixa mín–máx", {"bala": True}),
    ("Renas em 10 ms, 5 execuções por ponto", {"bala": True, "cor": CINZA, "tam": 14}),
], tam=16, espaco=8)
rodape(s, 7)
notas(s, "[Integrante 3] (~45 s) Primeiro gráfico: variamos o número de elfos com as renas chegando junto, "
         "5 execuções por ponto. A vazão cresce porque os 10 ms iniciais e os 5 ms de viagem são um custo "
         "fixo, que se dilui quando há mais grupos. A fila FIFO atende os grupos sem pausa. Os valores "
         "ficaram cerca de 2% abaixo da previsão analítica, por causa da latência dos sleeps, e a variação "
         "entre repetições é bem menor que a diferença entre os níveis.")

# ------------------------------------------------------------------ 8. G2
s = novo()
titulo(s, "O preço da prioridade: renas junto atrasam o 1º grupo")
imagem(s, G / "G2_p95_renas.png", 0.6, 1.4, w=8.6)
texto(s, 9.6, 1.6, 3.2, 1.1, f"{br(r06['espera_p95_ms_media_execucoes'], 1)} ms", tam=44, negrito=True,
      cor=VERMELHO, fonte=FONTE_T)
texto(s, 9.6, 2.55, 3.2, 0.7, "p95 com renas junto", tam=16, cor=CINZA)
texto(s, 9.6, 3.3, 3.2, 3.4, [
    (f"Antes e depois: {br(a05['espera_p95_ms_media_execucoes'], 1)} ms (só a formação do grupo)", {"bala": True}),
    ("Resposta às renas: dezenas de µs em todos os casos", {"bala": True}),
    ("Hipótese confirmada; ~3 ms ≫ faixa < 0,1 ms", {"bala": True}),
    ("6 elfos, 5 execuções por categoria", {"bala": True, "cor": CINZA, "tam": 14}),
], tam=16, espaco=8)
rodape(s, 8)
notas(s, "[Integrante 3] (~45 s) Segundo gráfico: 6 elfos e a chegada das renas variando. Quando as renas "
         "chegam junto com os elfos, a viagem ocupa Noel de 10 a 15 ms, justamente enquanto o primeiro grupo "
         "se forma, e o p95 sobe para cerca de 5 ms. Antes ou depois, Noel está livre e a espera é só o tempo "
         "de completar o grupo, uns 2 ms. Ao lado, as renas são atendidas em dezenas de microssegundos: esse "
         "é o efeito da prioridade.")

# ------------------------------------------------------------------ 9. G3 + G4
s = novo()
titulo(s, "Dentro da referência: o atraso se propaga sem crescer")
imagem(s, G / "G3_p95_grupo.png", 0.6, 1.4, w=5.0)
imagem(s, G / "G4_linha_tempo.png", 5.85, 1.45, w=6.9)
texto(s, 0.6, 5.0, 5.0, 1.9, [
    ("G3: grupos 1 e 2 com p95 ≈ 5,1–5,2 ms", {"negrito": True}),
    ("Os grupos chegam a cada 3 ms e cada ajuda dura 3 ms, então a fila não cresce", {"tam": 14}),
], tam=16, espaco=4)
texto(s, 5.85, 4.45, 6.9, 2.4, [
    ("G4: disputa real (E06_R10_rep1)", {"negrito": True}),
    ("Os dois grupos ficam prontos durante a preparação da viagem e esperam sem interrompê-la; "
     "depois Noel fica ocupado sem pausa", {"tam": 14}),
], tam=16, espaco=4)
rodape(s, 9)
notas(s, "[Integrante 3] (~45 s) No G3, os dois grupos da referência têm praticamente o mesmo p95. Esperávamos "
         "que o segundo esperasse mais, mas como os grupos chegam a cada 3 ms e cada ajuda leva 3 ms, o atraso "
         "causado pela viagem se propaga sem acumular. O G4 mostra essa execução no tempo: a 9ª rena chega "
         "em 10 ms, a viagem vai até 15 ms, os grupos esperam na faixa hachurada e depois são atendidos em sequência.")

# ------------------------------------------------------------------ 10. conclusões
s = novo(VERDE)
texto(s, 0.6, 0.35, 12.1, 0.9, "Conclusões e limitações", tam=34, cor=BRANCO, fonte=FONTE_T, negrito=True)
stats = [(f"{n_ok}/{n_exec}", "execuções válidas"), (str(viol), "violações na matriz"),
         (f"{int((est['status'] == 'ok').sum())}/{len(est)}", "execuções de estresse ok")]
for i, (v, rot) in enumerate(stats):
    x = 0.6 + i * 4.1
    texto(s, x, 1.45, 3.8, 1.1, v, tam=50, negrito=True, cor=BRANCO, fonte=FONTE_T)
    texto(s, x, 2.55, 3.8, 0.5, rot, tam=16, cor=BRANCO)
texto(s, 0.6, 3.35, 6.0, 3.5, [
    ("Conclusões", {"negrito": True, "tam": 20}),
    ("As regras valeram em todas as execuções: grupos de 3, renas primeiro, sem travamento", {"bala": True}),
    ("A vazão cresce com N; o p95 só sobe quando a viagem coincide com o 1º grupo", {"bala": True}),
    ("Numa versão contínua, a prioridade absoluta às renas pode causar starvation dos elfos (mitigação: aging/cota)", {"bala": True}),
], tam=16, cor=BRANCO, espaco=8)
texto(s, 6.85, 3.35, 5.9, 3.5, [
    ("Limitações", {"negrito": True, "tam": 20}),
    ("Cenário quase determinístico: variação pequena entre repetições", {"bala": True}),
    ("Latência de timers e do SO (0,05–0,5 ms) entra nas medidas", {"bala": True}),
    ("p95 de 3 ou 6 amostras = máximo; sem TSan/Helgrind disponíveis na máquina", {"bala": True}),
], tam=16, cor=BRANCO, espaco=8)
rodape(s, 10, escuro=True)
notas(s, "[Integrante 3] (~45 s) Fechando: 25 de 25 execuções válidas, nenhuma violação e 300 execuções "
         "de estresse sem travar. As duas hipóteses se confirmaram. Como limitação, o cenário é dirigido pelo "
         "tempo, então a variação é pequena, e a latência dos timers entra nas medidas. Numa versão com "
         "demandas contínuas, a prioridade absoluta às renas pode deixar os elfos esperando sem limite; "
         "a solução seria aging ou uma cota. Obrigado, perguntas?")

prs.save(RAIZ / "slides.pptx")
print("slides.pptx gerado")
