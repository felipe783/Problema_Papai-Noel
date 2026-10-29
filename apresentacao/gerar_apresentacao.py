#!/usr/bin/env python3
"""gerar_apresentacao.py — monta apresentacao_papai_noel.pptx (v4).

Deck novo, planejado com as skills de apresentação (pirâmide + SCQA, ghost deck,
títulos de ação, um acento por slide, page tracker). Mantém a identidade visual da
v3 (verde dominante, vermelho de acento, laranja = preparação e azul = ajuda,
como nos gráficos G1–G4).

Regras de honestidade:
- todo trecho de código é lido de ../src (nunca digitado à mão);
- números de resultados vêm de ../dados (tabela_esperas.csv, eventos.csv, resumo.csv);
- os cenários "renas durante a ajuda" (--renas-ms 13) e "empate natural"
  (--renas-ms 12, 30 execuções) são verificações feitas em 2026-10-28 fora da
  rodada oficial; os valores estão fixados abaixo e marcados como tal nos slides.

Uso: python3 apresentacao/gerar_apresentacao.py   (requer python-pptx e pandas)
"""
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
D, G, SRC = RAIZ / "dados", RAIZ / "graficos", RAIZ / "src"
SAIDA = AQUI / "apresentacao_papai_noel.pptx"

# ------------------------------------------------------------------ paleta (papéis, não enfeites)
TINTA = RGBColor(0x1A, 0x1A, 0x1A)       # texto principal
CINZA = RGBColor(0x55, 0x55, 0x55)       # texto secundário
CINZA_M = RGBColor(0x8A, 0x8C, 0x88)     # rótulos discretos
LINHA = RGBColor(0xD3, 0xD5, 0xD0)       # filetes
FUNDO2 = RGBColor(0xF3, 0xF4, 0xF2)      # regiões neutras
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERDE = RGBColor(0x1F, 0x4D, 0x3A)       # estrutura (títulos, cabeçalhos)
VERDE_CL = RGBColor(0xE6, 0xEF, 0xEA)
ACENTO = RGBColor(0xB3, 0x1B, 0x2B)      # o único destaque de cada slide
ACENTO_CL = RGBColor(0xF8, 0xE5, 0xE7)
# codificação categórica dos gráficos (mesmas cores do G4)
PREP = RGBColor(0xEB, 0x68, 0x34)        # preparação da viagem
AJUDA = RGBColor(0x2A, 0x78, 0xD6)       # ajuda a um grupo
RENA = RGBColor(0x5B, 0x4F, 0xD6)        # 9ª rena
PRONTO = RGBColor(0x1B, 0xAA, 0x74)      # grupo pronto
DORME = RGBColor(0xC9, 0xC8, 0xC3)       # dormindo / esperando
# código
COD_BG = RGBColor(0x1E, 0x22, 0x20)
COD_TX = RGBColor(0xE6, 0xE8, 0xE4)
COD_NUM = RGBColor(0x7E, 0x86, 0x80)
COD_DEST = RGBColor(0xF2, 0xC9, 0x4C)

F_TIT, F_TXT, F_COD = "Cambria", "Calibri", "Courier New"
SECOES = ["Fundamentos", "Problema", "Solução", "Sincronização", "Propriedades", "Resultados"]


def br(x, n=2):
    return f"{x:.{n}f}".replace(".", ",")


# ------------------------------------------------------------------ dados reais
tab = pd.read_csv(D / "tabela_esperas.csv").set_index("configuracao").to_dict("index")
res = pd.read_csv(D / "resumo.csv")
ev = pd.read_csv(D / "eventos.csv")
ev_tp = pd.read_csv(D / "teste_prioridade" / "eventos.csv")
ev_in = pd.read_csv(D / "ingenua" / "eventos.csv")
tp = pd.read_csv(D / "teste_prioridade" / "resumo.csv").iloc[0]
ing = pd.read_csv(D / "ingenua" / "resumo.csv").iloc[0]
est = pd.read_csv(D / "estresse" / "resumo.csv")
esp = pd.read_csv(D / "espera_elfos.csv")
N_EXEC, N_OK = len(res), int((res["status"] == "ok").sum())
VIOL = int(res["violacoes_prioridade"].sum() + res["violacoes_regras"].sum())
EST_OK, EST_N = int((est["status"] == "ok").sum()), len(est)
_g3 = esp[esp["execucao"].str.startswith("E06_R10")].groupby(["execucao", "grupo_id"])["espera_ms"].max().unstack()
G3 = {g: _g3[g].mean() for g in _g3.columns}
r03, r06, r12, a05, d30 = (tab[k] for k in ("E03_R10", "E06_R10", "E12_R10", "E06_R05", "E06_R30"))

# verificações fora da rodada oficial (2026-10-28, binário compilado de src/)
R13 = {"g1_pronto": 12.271, "dec1": 12.334, "nona": 13.116, "g2_pronto": 15.083, "fim1": 15.388,
       "dec2": 15.398, "fimprep": 20.470, "dec3": 20.485, "fim3": 23.558, "enc": 23.565}
EMPATE = {"viagem": 12, "grupo": 18, "n": 30}

COD = {f: (SRC / f).read_text().splitlines() for f in ("main.c", "noel.c", "elfo.c", "rena.c", "comum.h")}


def faixas(df, execucao):
    e = df[df["execucao"] == execucao]

    def par(a, b):
        ini, fim = e[e["evento"] == a], e[e["evento"] == b]
        return list(zip(ini["t_ms"], fim["t_ms"], ini["grupo_id"]))

    return {"prep": par("inicio_preparacao", "fim_preparacao"), "ajuda": par("inicio_ajuda", "fim_ajuda"),
            "prontos": list(e[e["evento"] == "grupo_pronto"]["t_ms"]),
            "nona": float(e[e["evento"] == "nona_rena"]["t_ms"].iloc[0]),
            "fim": float(e[e["evento"] == "encerramento"]["t_ms"].iloc[0])}


# ------------------------------------------------------------------ infraestrutura
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
LAYOUT = prs.slide_layouts[6]
NUM = [0]


def novo():
    s = prs.slides.add_slide(LAYOUT)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BRANCO
    NUM[0] += 1
    return s


def _sem_estilo(sh):
    """Remove o p:style do tema (sombra e contorno herdados): cor e linha ficam explícitas."""
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def _run(p, txt, tam, cor, fonte, negrito=False, italico=False):
    r = p.add_run()
    r.text = txt
    f = r.font
    f.size, f.name, f.bold, f.italic = Pt(tam), fonte, negrito, italico
    f.color.rgb = cor
    return r


def texto(s, x, y, w, h, pars, tam=16, cor=TINTA, fonte=F_TXT, negrito=False, alinh=PP_ALIGN.LEFT,
          ancora=MSO_ANCHOR.TOP, espaco=4, margem=0.04, entrelinha=None):
    """pars: str | lista. Cada item: str, ou lista de runs (txt, {opções}), ou (str, {opções do parágrafo})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(margem)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    if isinstance(pars, str):
        pars = [pars]
    for i, item in enumerate(pars):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinh
        p.space_after = Pt(espaco)
        if entrelinha:
            p.line_spacing = entrelinha
        popt = {}
        if isinstance(item, tuple):
            item, popt = item
        runs = item if isinstance(item, list) else [(item, {})]
        for txt, o in runs:
            o = {**popt, **o}
            _run(p, txt, o.get("tam", tam), o.get("cor", cor), o.get("fonte", fonte),
                 o.get("negrito", negrito), o.get("italico", False))
    return tb


def caixa(s, x, y, w, h, fundo=None, borda=None, esp=0.75, forma=MSO_SHAPE.RECTANGLE, txt=None, tam=14,
          cor=TINTA, negrito=False, alinh=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE, fonte=F_TXT, margem=0.08):
    sh = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.08
    _sem_estilo(sh)
    if fundo is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fundo
    if borda is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borda
        sh.line.width = Pt(esp)
    if txt is not None:
        tf = sh.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = ancora
        tf.margin_left = tf.margin_right = Inches(margem)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        linhas = txt if isinstance(txt, list) else [txt]
        for i, item in enumerate(linhas):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = alinh
            runs = item if isinstance(item, list) else [(item, {})]
            for t, o in runs:
                _run(p, t, o.get("tam", tam), o.get("cor", cor), o.get("fonte", fonte), o.get("negrito", negrito))
    return sh


def seta(s, x1, y1, x2, y2, cor=CINZA, esp=1.25, ponta=True, tracejado=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _sem_estilo(c)
    c.line.color.rgb = cor
    c.line.width = Pt(esp)
    ln = c.line._get_or_add_ln()
    if tracejado:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    if ponta:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("len", "med")
    return c


def titulo(s, txt, secao=None):
    if secao is not None:
        tracker(s, secao)
    texto(s, 0.6, 0.62, 12.1, 0.95, txt, tam=25, cor=VERDE, fonte=F_TIT, negrito=True,
          ancora=MSO_ANCHOR.MIDDLE, espaco=0)
    seta(s, 0.6, 1.62, 12.73, 1.62, cor=LINHA, esp=0.75, ponta=False)


def tracker(s, atual):
    """Page tracker: seções na faixa superior, a atual em verde com sublinhado."""
    nomes = SECOES + ["Apoio"] if atual == "Apoio" else SECOES
    x = 0.6
    for nome in nomes:
        ativo = nome == atual
        w = 0.09 * len(nome) + 0.5
        texto(s, x, 0.2, w, 0.3, nome, tam=11, cor=VERDE if ativo else CINZA_M, negrito=ativo, espaco=0)
        if ativo:
            caixa(s, x + 0.04, 0.48, w - 0.4, 0.035, fundo=VERDE)
        x += w


def rodape(s, fonte=None):
    if fonte:
        texto(s, 0.6, 7.03, 10.5, 0.3, f"Fonte: {fonte}", tam=10, cor=CINZA_M, espaco=0)
    texto(s, 12.0, 7.03, 0.73, 0.3, str(NUM[0]), tam=10, cor=CINZA_M, alinh=PP_ALIGN.RIGHT, espaco=0)


def notas(s, txt):
    s.notes_slide.notes_text_frame.text = txt


def codigo(s, x, y, w, arq, trechos, dest=(), tam=12, titulo_cod=None):
    """Bloco de código com o texto real de src/arq. trechos: [(ini, fim), ...] (base 1)."""
    linhas = []
    for k, (a, b) in enumerate(trechos):
        if k:
            linhas.append((None, "…"))
        linhas += [(n, COD[arq][n - 1]) for n in range(a, b + 1)]
    reais = [t for n, t in linhas if n is not None and t.strip()]
    corte = min(len(t) - len(t.lstrip()) for t in reais)
    alt_l = tam * 1.2 / 72
    maxc = int((w - 0.16) / (tam * 0.6 / 72)) - 6
    h = len(linhas) * alt_l + 0.2 + (0.32 if titulo_cod else 0)
    caixa(s, x, y, w, h, fundo=COD_BG)
    y0 = y + 0.08
    if titulo_cod:
        texto(s, x + 0.1, y0, w - 0.2, 0.3, titulo_cod, tam=11, cor=COD_NUM, fonte=F_COD, espaco=0)
        y0 += 0.32
    tb = s.shapes.add_textbox(Inches(x + 0.08), Inches(y0), Inches(w - 0.16), Inches(h - 0.16))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (n, t) in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(0)
        p.line_spacing = 1.0
        if n is None:
            _run(p, "     " + t, tam, COD_NUM, F_COD)
            continue
        d = n in dest
        _run(p, f"{n:>3}  ", tam, COD_DEST if d else COD_NUM, F_COD, negrito=d)
        t = t[corte:]
        if len(t) > maxc:
            t = t[:maxc - 1].rstrip() + "…"
        _run(p, t, tam, COD_DEST if d else COD_TX, F_COD, negrito=d)
    return y + h


def tabela(s, x, y, larguras, linhas, tam=13, alt=0.36, cab=True, destaque_linhas=(), destaque_cor=ACENTO_CL,
           negrito_col0=False):
    nl, nc = len(linhas), len(larguras)
    gf = s.shapes.add_table(nl, nc, Inches(x), Inches(y), Inches(sum(larguras)), Inches(alt * nl))
    t = gf.table
    tblPr = t._tbl.tblPr
    for a in ("bandRow", "firstRow"):
        tblPr.set(a, "0")
    st = tblPr.find(qn("a:tableStyleId"))
    if st is not None:
        st.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
    for j, w in enumerate(larguras):
        t.columns[j].width = Inches(w)
    for i, row in enumerate(linhas):
        t.rows[i].height = Inches(alt)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.07)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if cab and i == 0:
                c.fill.fore_color.rgb = VERDE
            elif i in destaque_linhas:
                c.fill.fore_color.rgb = destaque_cor
            else:
                c.fill.fore_color.rgb = BRANCO if i % 2 else FUNDO2
            tf = c.text_frame
            tf.word_wrap = True
            runs = val if isinstance(val, list) else [(str(val), {})]
            p = tf.paragraphs[0]
            for txt, o in runs:
                eh_cab = cab and i == 0
                _run(p, txt, o.get("tam", tam), BRANCO if eh_cab else o.get("cor", TINTA),
                     o.get("fonte", F_TXT), eh_cab or o.get("negrito", negrito_col0 and j == 0))
    return gf


def imagem(s, arq, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(arq), Inches(x), Inches(y), **kw)


def linha_tempo(s, x, y, w, t0, t1, linhas, alt=0.42, gap=0.2, ticks=None, rot_w=1.5, tam=11):
    """linhas: [{rot, segs:[(a,b,cor,txt)], marcas:[(t,cor,txt,nivel)]}]; nivel 0 = rótulo abaixo, 1-2 = acima. Devolve y final."""
    eixo_x, eixo_w = x + rot_w, w - rot_w
    px = lambda t: eixo_x + (t - t0) / (t1 - t0) * eixo_w
    yy = y
    for ln in linhas:
        texto(s, x, yy, rot_w - 0.1, alt, ln["rot"], tam=12, cor=TINTA, negrito=ln.get("neg", False),
              ancora=MSO_ANCHOR.MIDDLE, alinh=PP_ALIGN.RIGHT, espaco=0)
        caixa(s, eixo_x, yy + alt / 2 - 0.01, eixo_w, 0.02, fundo=LINHA)
        for t, cor, txt, nivel in ln.get("marcas", []):
            caixa(s, px(t) - 0.012, yy - 0.08, 0.024, alt + 0.16, fundo=cor)
            if txt:
                ty = yy - 0.08 - 0.27 * int(nivel) if nivel else yy + alt + 0.08
                texto(s, px(t) - 1.0, ty - (0.0 if not nivel else 0.0), 2.0, 0.26, txt, tam=10, cor=cor,
                      alinh=PP_ALIGN.CENTER, espaco=0, negrito=True)
        for a, b, cor, txt in ln.get("segs", []):
            caixa(s, px(a), yy + 0.04, max(px(b) - px(a), 0.03), alt - 0.08, fundo=cor,
                  txt=txt, tam=tam, cor=BRANCO, negrito=True, margem=0.02)
        yy += alt + gap
    for tk in (ticks or []):
        caixa(s, px(tk), yy - gap + 0.02, 0.01, 0.07, fundo=CINZA_M)
        texto(s, px(tk) - 0.4, yy - gap + 0.08, 0.8, 0.25, br(tk, 0) if tk == int(tk) else br(tk, 1),
              tam=10, cor=CINZA_M, alinh=PP_ALIGN.CENTER, espaco=0)
    return yy, px


def legenda_cores(s, x, y, itens, tam=11):
    for cor, txt in itens:
        caixa(s, x, y + 0.06, 0.22, 0.16, fundo=cor)
        texto(s, x + 0.28, y, 2.6, 0.28, txt, tam=tam, cor=CINZA, espaco=0)
        x += 0.28 + 0.085 * len(txt) + 0.35


def rotulo(s, x, y, w, txt, cor=VERDE, tam=12):
    """Rótulo de seção dentro do slide (versalete discreto)."""
    texto(s, x, y, w, 0.3, txt.upper(), tam=tam, cor=cor, negrito=True, espaco=0)


def callout(s, x, y, w, h, txt, tam=14):
    """Faixa de conclusão com o acento: um por slide."""
    caixa(s, x, y, 0.07, h, fundo=ACENTO)
    texto(s, x + 0.18, y, w - 0.18, h, txt, tam=tam, cor=TINTA, ancora=MSO_ANCHOR.MIDDLE, espaco=0)


# ================================================================== SLIDES
# 1 — capa ----------------------------------------------------------
s = novo()
caixa(s, 0, 0, 0.35, 7.5, fundo=VERDE)
texto(s, 1.0, 1.0, 11, 0.4, "SISTEMAS OPERACIONAIS · TEMA 11", tam=14, cor=VERDE, negrito=True)
texto(s, 1.0, 1.6, 11.3, 1.2, "Problema do Papai Noel", tam=48, cor=TINTA, fonte=F_TIT, negrito=True)
texto(s, 1.0, 2.75, 11.0, 1.2,
      "Como um mutex e quatro variáveis de condição coordenam Noel, 9 renas e N elfos em C11 com pthreads",
      tam=22, cor=CINZA, fonte=F_TIT)
caixa(s, 1.0, 4.25, 2.2, 0.04, fundo=ACENTO)
texto(s, 1.0, 4.5, 11, 1.4, [
    [("Integrante: ", {"negrito": True}), ("Felipe Prior Mulinari", {})],
    [("Professor: ", {"negrito": True}), ("Lucca Abbado Neres · FAG", {})],
    [("Implementação: ", {"negrito": True}), ("tema11_papai_noel/src (C11 + POSIX Threads, GCC 15.3.1, Fedora 43)", {})],
], tam=16, cor=TINTA, espaco=8)
notas(s, "Abertura (≈20 s). Apresentar-se e dizer a mensagem central em uma frase: "
         "um mutex protege todo o estado, quatro variáveis de condição fazem as threads esperarem sem gastar CPU, "
         "e só o Noel decide, testando as renas primeiro.")

# 2 — resumo executivo (SCQA + pirâmide) ----------------------------
s = novo()
titulo(s, "Um mutex, quatro variáveis de condição e uma regra de decisão garantem as regras do Papai Noel")
texto(s, 0.6, 1.8, 12.1, 0.75, [[
    ("Situação: ", {"negrito": True, "cor": VERDE}), ("threads compartilham memória e o SO escolhe a ordem em que rodam.  ", {}),
    ("Complicação: ", {"negrito": True, "cor": VERDE}), ("três regras do problema dependem dessa ordem.  ", {}),
    ("Pergunta: ", {"negrito": True, "cor": VERDE}), ("como coordená-las sem corrida, sem deadlock e sem espera ocupada?", {}),
]], tam=15, cor=TINTA)
razoes = [
    ("1", "O mutex torna “testar e alterar” indivisível",
     "O 3º elfo fecha o grupo sem soltar o mutex: um 4º elfo nunca entra num grupo fechado.", "elfo.c:41-53"),
    ("2", "As variáveis de condição fazem esperar sem gastar CPU",
     "Noel, elfos e renas dormem até um predicado valer e são acordados por quem o tornou verdadeiro.", "4 × pthread_cond_t"),
    ("3", "Só Noel decide, com o mutex, testando as renas primeiro",
     "A prioridade vale entre tarefas: uma ajuda iniciada nunca é interrompida.", "noel.c:89-90, 117"),
]
for i, (n, cab, corpo, ref) in enumerate(razoes):
    x = 0.6 + i * 4.1
    caixa(s, x, 2.75, 3.9, 2.75, fundo=FUNDO2)
    texto(s, x + 0.2, 2.9, 0.6, 0.6, n, tam=30, cor=VERDE, fonte=F_TIT, negrito=True)
    texto(s, x + 0.2, 3.5, 3.5, 0.8, cab, tam=16, cor=TINTA, negrito=True)
    texto(s, x + 0.2, 4.3, 3.5, 0.9, corpo, tam=13, cor=CINZA)
    texto(s, x + 0.2, 5.12, 3.5, 0.3, ref, tam=11, cor=VERDE, fonte=F_COD)
callout(s, 0.6, 5.8, 12.1, 0.85,
        [[("Evidência: ", {"negrito": True}),
          (f"{N_OK}/{N_EXEC} execuções válidas e {VIOL} violações · teste de prioridade PASSOU · a versão ingênua gerou "
           f"{int(ing['violacoes_prioridade'])} violações, detectadas · estresse {EST_OK}/{EST_N} ok", {})]], tam=15)
rodape(s, "src/; dados/resumo.csv; dados/log_execucao.txt")
notas(s, "Resumo executivo (≈40 s). Ler a linha SCQA e depois as três razões. A evidência embaixo é o que "
         "voltaremos a ver na seção de resultados. Transição: 'Para chegar lá, primeiro os fundamentos.'")

# 3 — SO e threads ---------------------------------------------------
s = novo()
titulo(s, "O SO decide a ordem em que as threads rodam; sincronizar é impor as ordens que importam", "Fundamentos")
# processo com threads
caixa(s, 0.6, 1.95, 5.6, 4.3, borda=VERDE, esp=1.25)
texto(s, 0.8, 2.0, 5.2, 0.35, "1 processo: ./simulacao", tam=14, cor=VERDE, negrito=True)
caixa(s, 0.9, 2.45, 5.0, 1.0, fundo=VERDE_CL,
      txt=[[("Memória compartilhada", {"negrito": True})], [("struct Oficina: contadores, grupos, estado de Noel", {"tam": 12})]],
      tam=14)
for i, (nome, q) in enumerate([("Noel", "1"), ("Renas", "9"), ("Elfos", "N")]):
    xx = 0.9 + i * 1.7
    caixa(s, xx, 4.05, 1.6, 1.0, fundo=BRANCO, borda=CINZA_M, txt=[[(nome, {"negrito": True})], [(f"{q} thread(s)", {"tam": 12})]])
    seta(s, xx + 0.8, 4.05, xx + 0.8, 3.48, cor=CINZA_M)
texto(s, 0.9, 5.2, 5.1, 0.95, "Cada thread tem a própria pilha, mas lê e escreve a mesma Oficina. "
      "O escalonador do SO pode pausar qualquer uma entre duas instruções.", tam=13, cor=CINZA)
# definições
defs = [
    ("Concorrência", "várias threads em andamento ao mesmo tempo, intercaladas ou em núcleos diferentes."),
    ("Seção crítica", "trecho que lê ou altera dados compartilhados; só uma thread por vez pode executá-lo."),
    ("Condição de corrida", "o resultado depende da ordem do entrelaçamento, que muda a cada execução."),
    ("Sincronização", "exclusão mútua nos dados + uma forma de esperar e avisar eventos sem gastar CPU."),
]
for i, (k, v) in enumerate(defs):
    yy = 1.95 + i * 1.08
    texto(s, 6.6, yy, 6.1, 1.0, [[(k, {"negrito": True, "cor": ACENTO if i == 3 else VERDE, "tam": 16})], v],
          tam=14, espaco=2)
rodape(s, "src/comum.h:79-123 (struct Oficina); src/main.c:187-204 (criação das threads)")
notas(s, "Q1 e Q2 do professor. Processo = espaço de memória próprio; thread = fluxo de execução dentro do processo, "
         "compartilhando a memória. O SO cria e escalona as threads; como não controlamos a ordem, precisamos de "
         "mecanismos. Quando uma thread bloqueia num mutex ou cond var, o SO a tira da CPU (no Linux, via futex).")

# 4 — if + contador -----------------------------------------------------
s = novo()
titulo(s, "Um if com contador falha porque testar e alterar são passos que outra thread pode intercalar", "Fundamentos")
rotulo(s, 0.6, 1.85, 6, "Sem mutex: renas_prontas++ vira ler → somar → gravar", cor=ACENTO)
tabela(s, 0.6, 2.25, [0.6, 2.0, 2.0, 1.2], [
    ["t", "Rena A", "Rena B", "memória"],
    ["1", "lê 7", "", "7"],
    ["2", "", "lê 7", "7"],
    ["3", "grava 8", "", "8"],
    ["4", "", [("grava 8", {"negrito": True, "cor": ACENTO})], [("8 ✗", {"negrito": True, "cor": ACENTO})]],
], tam=14, alt=0.4)
texto(s, 0.6, 4.35, 5.8, 0.9, "As duas renas chegaram, mas o contador só subiu 1. Com 9 renas presentes ele "
      "para em 8 e Noel nunca prepara a viagem.", tam=14, cor=TINTA)
rotulo(s, 6.9, 1.85, 6, "Com mutex: o par vira indivisível")
tabela(s, 6.9, 2.25, [0.6, 2.4, 1.9, 0.9], [
    ["t", "Rena A", "Rena B", "mem."],
    ["1", "lock · lê 7 · grava 8", "lock → bloqueia", "8"],
    ["2", "unlock", "(acorda)", "8"],
    ["3", "", "lê 8 · grava 9 · unlock", [("9 ✓", {"negrito": True, "cor": VERDE})]],
], tam=14, alt=0.4)
fim4 = codigo(s, 6.9, 4.0, 5.8, "rena.c", [(23, 23), (28, 28)], dest={28}, tam=13)
texto(s, 6.9, fim4 + 0.05, 5.8, 0.35, "Mesma seção crítica: lock na linha 23, ++ na 28.", tam=12, cor=CINZA)
callout(s, 0.6, 5.55, 12.1, 1.15,
        [[("O mesmo vale para o elfo: ", {"negrito": True}),
          ("sem mutex, dois elfos leem n == 2, os dois passam no if e escrevem em membros[2], formando um “grupo” de 4. "
           "Um lock caseiro com flag (while (ocupado); ocupado = 1;) sofre a mesma corrida e ainda gasta CPU.", {})]], tam=14)
rodape(s, "src/rena.c:23-28; src/elfo.c:34-41")
notas(s, "Q4. Mensagem: o problema não é o if, é o intervalo entre o teste e a alteração. O mutex não deixa o "
         "teste-e-alteração ser intercalado. No código, o if do elfo existe, mas está DENTRO do mutex.")

# 5 — mutex × semáforo × cond var ----------------------------------------
s = novo()
titulo(s, "Mutex dá exclusão, semáforo conta e variável de condição espera um predicado: usamos mutex e condição",
       "Fundamentos")
tabela(s, 0.6, 1.9, [2.2, 3.3, 3.3, 3.3], [
    ["", "Mutex", "Semáforo", "Variável de condição"],
    ["Para quê", "exclusão mútua na seção crítica", "contar recursos ou avisos", "esperar até um predicado valer"],
    ["Operações", "lock / unlock", "wait (P) / post (V)", "wait / signal / broadcast"],
    ["Valor inicial", "nenhum: começa destravado", "um número: sem_init(&s, 0, v)", "nenhum: fila vazia"],
    ["Tem dono?", "sim: só quem travou destrava", "não: qualquer thread dá post", "não; sempre usada com um mutex"],
    ["Lembra avisos?", "—", "sim, no contador", "não: por isso exige o predicado"],
    [[("No projeto", {"negrito": True})], [("1 mutex (main.c:131)", {"negrito": True})],
     [("nenhum (0 ocorrências de sem_)", {"negrito": True})], [("4 cond vars (main.c:132-139)", {"negrito": True})]],
], tam=14, alt=0.55, negrito_col0=True, destaque_linhas=(6,))
callout(s, 0.6, 5.95, 12.1, 0.75,
        "Por que não semáforos: Noel precisa ler as duas demandas ao mesmo tempo para escolher. Com mutex + cond var, "
        "ele lê renas_ok e elfos_ok na mesma seção crítica e decide.", tam=14)
rodape(s, "src/comum.h:82-86; src/main.c:131-139; grep -c sem_ src/* = 0")
notas(s, "Q3 e Q10. Semáforo binário imita um mutex, mas sem dono. Variável de condição não guarda aviso: se ninguém "
         "está esperando, o signal se perde, e por isso a thread sempre testa o estado antes de dormir.")

# 6 — o problema -----------------------------------------------------------
s = novo()
titulo(s, "Noel atende duas demandas, e três das regras dependem da ordem entre as threads", "Problema")
caixa(s, 0.6, 2.2, 1.9, 0.9, fundo=FUNDO2, txt=[[("N elfos", {"negrito": True})], [("elfo i chega em 10 + i ms", {"tam": 11})]])
caixa(s, 2.9, 2.2, 1.9, 0.9, fundo=FUNDO2, txt=[[("Grupos de 3", {"negrito": True})], [("fila FIFO", {"tam": 11})]])
caixa(s, 2.75, 4.2, 2.2, 0.9, fundo=VERDE, cor=BRANCO, txt=[[("Papai Noel", {"negrito": True})], [("1 thread, decide", {"tam": 11})]])
caixa(s, 0.6, 5.6, 1.9, 0.9, fundo=FUNDO2, txt=[[("9 renas", {"negrito": True})], [("chegam juntas em R ms", {"tam": 11})]])
seta(s, 2.5, 2.65, 2.9, 2.65)
seta(s, 3.85, 3.1, 3.85, 4.2)
seta(s, 1.55, 5.6, 2.75, 4.95)
texto(s, 3.95, 3.4, 1.6, 0.6, "ajuda:\n3 ms por grupo", tam=12, cor=AJUDA, negrito=True)
texto(s, 2.75, 5.25, 2.4, 0.6, "preparação da viagem:\n5 ms", tam=12, cor=PREP, negrito=True)
regras = [
    (False, "Noel dorme até haver um grupo de 3 elfos ou as 9 renas."),
    (True, "Cada ajuda tem exatamente 3 elfos; um 4º não entra num grupo já fechado."),
    (True, "A viagem só é preparada depois que as 9 renas chegaram."),
    (True, "Com as duas demandas prontas na decisão, renas primeiro, sem interromper uma ajuda já iniciada."),
    (False, "Versão finita: 1 viagem e 1 pedido por elfo; depois da viagem, os grupos restantes."),
    (False, "Fim: todos encerram quando as duas demandas terminam."),
]
rotulo(s, 5.6, 1.9, 7, "Regras desta versão")
for i, (crit, r) in enumerate(regras):
    yy = 2.3 + i * 0.68
    caixa(s, 5.6, yy + 0.08, 0.32, 0.32, fundo=ACENTO if crit else LINHA, forma=MSO_SHAPE.OVAL,
          txt=str(i + 1), tam=11, cor=BRANCO if crit else TINTA, negrito=True, margem=0)
    texto(s, 6.05, yy, 6.65, 0.62, r, tam=15, cor=TINTA, negrito=crit, ancora=MSO_ANCHOR.MIDDLE, espaco=0)
texto(s, 5.6, 6.4, 7.1, 0.4, "● em vermelho: regras que só valem com sincronização correta", tam=12, cor=ACENTO)
rodape(s, "enunciado do Tema 11; README.md; src/main.c (padrões --elfo0-ms 10, --ajuda-ms 3, --preparo-ms 5)")
notas(s, "Explicar a história em 30 s. Destacar as três regras vermelhas: cada uma vai ganhar um slide na seção de "
         "sincronização (13, 15-16). Durações são simuladas com sleeps; nenhuma coordenação depende delas.")

# 7 — analogia com SO --------------------------------------------------------
s = novo()
titulo(s, "O problema é um escalonador em miniatura: um recurso único, clientes com prioridade, sem preempção",
       "Problema")
tabela(s, 0.6, 1.9, [4.0, 4.4, 3.7], [
    ["No Papai Noel", "No Sistema Operacional", "No código"],
    ["Noel: um só, disputado", "CPU ou dispositivo único", "thread_noel"],
    ["Elfos e renas", "processos/threads clientes", "thread_elfo, thread_rena"],
    ["Grupo de 3 elfos", "lote de requisições (amortiza custo)", "fila de grupos com id"],
    ["Renas antes dos elfos", "escalonamento por prioridade", "escolhe_renas = renas_ok"],
    ["Não interromper uma ajuda", "escalonamento não preemptivo", "decide só no topo do laço"],
    ["Noel dormindo", "thread bloqueada, fora da fila de prontos", "pthread_cond_timedwait"],
    ["Renas voltando sempre", "starvation → aging", "fora do escopo (versão finita)"],
], tam=14, alt=0.5, negrito_col0=True)
callout(s, 0.6, 6.05, 12.1, 0.7,
        "Limite da analogia: Noel não é o escalonador do SO. O SO continua escalonando as 1 + 9 + N threads; "
        "Noel é só uma delas, que escolhe a próxima tarefa.", tam=14)
rodape(s, "src/noel.c:76-147; vault: Perguntas da Apresentação, Q5")
notas(s, "Q5. Se o professor perguntar se Noel é preemptivo: não, a decisão só acontece entre tarefas, no topo do laço "
         "(noel.c:85). Threads × processos: os conceitos são os mesmos; entre processos seria preciso memória "
         "compartilhada e atributo PTHREAD_PROCESS_SHARED. Aqui o mutex é privado ao processo.")

# 8 — arquitetura ------------------------------------------------------------
s = novo()
titulo(s, "Cada personagem é uma thread, e todas compartilham uma única struct Oficina protegida pelo mutex",
       "Solução")
caixa(s, 0.6, 2.0, 1.6, 0.8, fundo=FUNDO2, txt=[[("main", {"negrito": True})], [("main.c", {"tam": 11, "fonte": F_COD})]])
for i, (nome, arq) in enumerate([("Noel ×1", "noel.c"), ("Rena ×9", "rena.c"), ("Elfo ×N", "elfo.c")]):
    yy = 1.95 + i * 1.25
    caixa(s, 3.0, yy, 1.8, 0.85, fundo=BRANCO, borda=CINZA_M,
          txt=[[(nome, {"negrito": True})], [(arq, {"tam": 11, "fonte": F_COD})]])
    seta(s, 2.2, 2.4, 3.0, yy + 0.42, cor=CINZA_M)
    seta(s, 4.8, yy + 0.42, 5.6, 3.35, cor=CINZA_M)
texto(s, 0.6, 2.9, 2.3, 0.6, "pthread_create\n(main.c:187-204)", tam=11, cor=CINZA)
caixa(s, 5.6, 1.95, 3.0, 4.05, fundo=VERDE_CL, borda=VERDE, esp=1.25)
texto(s, 5.75, 2.0, 2.8, 0.4, "struct Oficina", tam=15, cor=VERDE, negrito=True, fonte=F_COD)
campos = [("mutex + 4 cond vars", True), ("largada, threads_prontas, t0", False), ("grupos[], grupo_aberto,", False),
          ("  grupos_fechados, proximo_grupo", False), ("renas_prontas, viagem_*", False), ("estado de Noel, abortar", False),
          ("invariantes e eventos", False)]
texto(s, 5.75, 2.45, 2.8, 3.4, [[(c, {"negrito": b, "cor": ACENTO if b else TINTA})] for c, b in campos],
      tam=12, fonte=F_COD, espaco=5)
tabela(s, 8.95, 1.95, [1.45, 2.35], [
    ["Arquivo", "Responsabilidade"],
    ["main.c", "argumentos, init, threads, largada, joins"],
    ["noel.c", "laço de decisão, ajuda, viagem"],
    ["elfo.c", "chega, entra no grupo, espera"],
    ["rena.c", "chega, conta-se, espera"],
    ["tempo.c", "CLOCK_MONOTONIC e sleeps"],
    ["metricas.c, csv.c", "métricas e CSVs após os joins"],
], tam=12, alt=0.55)
texto(s, 0.6, 6.25, 12.1, 0.5, "Todo campo mutável da Oficina só é lido ou escrito com o mutex travado. "
      "As exceções são cfg e t0, escritos antes da largada e nunca mais alterados.", tam=13, cor=CINZA)
rodape(s, "src/comum.h:79-123; src/main.c:187-204")
notas(s, "Basta saber a estrutura em alto nível. Ponto-chave: há um único estado compartilhado e um único mutex. "
         "Isso simplifica o raciocínio e elimina deadlock por ordem de locks, ao custo de serializar o acesso.")

# 9 — fases da execução -------------------------------------------------------
s = novo()
titulo(s, "A execução tem quatro fases, e só a terceira aplica a política do problema", "Solução")
fases = [
    ("1 · Preparar", ["memset zera a Oficina", "init do mutex e das 4 cond vars", "cria 1 + 9 + N threads"],
     "main.c:109-140, 187-204"),
    ("2 · Largar", ["cada thread: prontas++ e espera", "main: espera todas, grava t0", "largada = 1 + broadcast"],
     "main.c:25-31, 207-213"),
    ("3 · Coordenar", ["elfos formam grupos de 3", "renas se contam", "Noel decide e executa tarefas"],
     "elfo.c · rena.c · noel.c"),
    ("4 · Encerrar", ["Noel sai do laço + broadcasts", "joins de todas as threads", "destroy, métricas, CSVs"],
     "noel.c:135-145; main.c:215-222"),
]
for i, (nome, itens, ref) in enumerate(fases):
    x = 0.6 + i * 3.08
    ativo = i == 2
    caixa(s, x, 2.0, 2.85, 0.65, fundo=ACENTO if ativo else VERDE, forma=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
          txt=nome, tam=16, cor=BRANCO, negrito=True)
    texto(s, x + 0.1, 2.9, 2.7, 1.6, [f"• {t}" for t in itens], tam=14, espaco=6)
    texto(s, x + 0.1, 4.45, 2.7, 0.5, ref, tam=11, cor=VERDE, fonte=F_COD)
caixa(s, 0.6, 5.2, 12.1, 0.02, fundo=LINHA)
texto(s, 0.6, 5.35, 12.1, 1.3, [
    [("Por que a largada comum: ", {"negrito": True}),
     ("todas as threads partem do mesmo t0, então os instantes medidos são comparáveis. É uma barreira feita à mão "
      "com contador + cond var, não pthread_barrier.", {})],
    [("Os sleeps só simulam durações ", {"negrito": True}),
     ("(clock_nanosleep absoluto: chegada, ajuda de 3 ms, preparação de 5 ms). Nenhuma coordenação depende deles.", {})],
], tam=14, espaco=8)
rodape(s, "src/main.c; src/tempo.c:43-50")
notas(s, "Fluxo de ponta a ponta. A fase 3 é o coração e ocupa os próximos slides. Na largada, as threads fazem "
         "broadcast (não signal) porque na cv_largada há outras threads esperando além da main: um signal poderia "
         "acordar a thread errada e a main nunca acordaria.")

# 10 — inventário dos mecanismos ------------------------------------------------
s = novo()
titulo(s, "Um mutex e quatro variáveis de condição, nenhum semáforo, coordenam o programa inteiro", "Sincronização")
tabela(s, 0.6, 1.85, [1.55, 2.0, 1.85, 2.35, 2.55, 1.8], [
    ["Mecanismo", "Inicializado em", "Começa", "Quem espera", "Quem avisa", "Coordena"],
    [[("mutex", {"fonte": F_COD, "negrito": True, "cor": ACENTO})], "main.c:131", "destravado, sem dono",
     "todos (lock)", "o dono (unlock)", "todo o estado"],
    [[("cv_largada", {"fonte": F_COD})], "main.c:132", "fila vazia", "threads (main.c:30); main (209)",
     "broadcast main.c:28, 212", "sinal de início"],
    [[("cv_noel", {"fonte": F_COD})], "main.c:136-139, CLOCK_MONOTONIC", "fila vazia", "só Noel: timedwait noel.c:101",
     "signal: 9ª rena rena.c:33; 3º elfo elfo.c:53", "há trabalho para Noel"],
    [[("cv_elfos", {"fonte": F_COD})], "main.c:133", "fila vazia", "elfos (elfo.c:58)",
     "broadcast noel.c:73, 143", "grupo concluído"],
    [[("cv_renas", {"fonte": F_COD})], "main.c:134", "fila vazia", "renas (rena.c:37)",
     "broadcast noel.c:41, 144", "viagem concluída"],
], tam=13, alt=0.68)
callout(s, 0.6, 6.05, 12.1, 0.7,
        "Variável de condição não tem valor: o que tem valor inicial são os predicados que ela observa, todos "
        "zerados pelo memset (main.c:109).", tam=14)
rodape(s, "src/comum.h:82-86; src/main.c:131-140; grep sem_|pthread_barrier em src/ = 0")
notas(s, "Q10 e Q11. Este é o slide de referência para 'quais mecanismos e onde'. signal só na cv_noel, porque ela "
         "tem um único consumidor (Noel). broadcast nas outras porque há várias threads esperando, e cada uma "
         "reavalia o próprio predicado.")

# 11 — inicialização ----------------------------------------------------------------
s = novo()
titulo(s, "O mutex nasce destravado em main.c:131, antes de qualquer thread existir", "Sincronização")
fim = codigo(s, 0.6, 1.85, 6.9, "main.c", [(107, 109), (131, 140)], dest={109, 131, 138, 139}, tam=13,
             titulo_cod="src/main.c · iniciar_oficina")
anot = [
    ("1", "memset (109)", "zera todos os predicados: 0 renas prontas, 0 grupos fechados, largada = 0, viagem não iniciada."),
    ("2", "pthread_mutex_init(&o->mutex, NULL) (131)",
     "NULL = atributos padrão: não recursivo, privado ao processo. Começa destravado e sem dono."),
    ("3", "setclock + cond_init (138-139)",
     "cv_noel usa CLOCK_MONOTONIC: o prazo do timedwait usa o mesmo relógio de t0 e não muda se a hora do sistema for ajustada."),
]
for i, (n, cab, corpo) in enumerate(anot):
    yy = 1.9 + i * 1.4
    caixa(s, 7.8, yy, 0.42, 0.42, fundo=ACENTO if i == 1 else VERDE, forma=MSO_SHAPE.OVAL, txt=n, tam=14,
          cor=BRANCO, negrito=True, margem=0)
    texto(s, 8.35, yy - 0.05, 4.4, 1.35, [[(cab, {"negrito": True, "fonte": F_COD, "tam": 12})], corpo], tam=13, espaco=2)
caixa(s, 0.6, 6.05, 12.1, 0.65, fundo=FUNDO2,
      txt=[[("Ordem no tempo:  ", {"negrito": True}), ("init (131)  →  pthread_create (187)  →  …  →  joins (215-219)  →  "
                                                      "destroy (222)", {"fonte": F_COD, "tam": 13})]],
      tam=14, alinh=PP_ALIGN.LEFT, margem=0.2)
rodape(s, "src/main.c:107-140, 150, 222")
notas(s, "Q9: 'como o mutex começa e qual linha o inicializa'. Resposta: main.c:131, destravado, atributos padrão. "
         "O memset NÃO conta como inicialização do mutex pela POSIX; o que vale é a linha 131. Não usamos "
         "PTHREAD_MUTEX_INITIALIZER. Ponto fraco honesto: os retornos dessas chamadas não são verificados.")

# 12 — padrão de espera -------------------------------------------------------------
s = novo()
titulo(s, "Toda espera segue o mesmo padrão: testar em while e dormir soltando o mutex", "Sincronização")
passos = [
    ("lock(mutex)", "entra na seção crítica", FUNDO2, TINTA),
    ("while (!predicado)", "o estado ainda não permite seguir?", FUNDO2, TINTA),
    ("cond_wait(&cv, &mutex)", "solta o mutex e dorme, de forma atômica", ACENTO, BRANCO),
    ("outra thread muda o estado", "e chama signal / broadcast", FUNDO2, TINTA),
    ("acorda e retoma o mutex", "volta ao while e testa de novo", FUNDO2, TINTA),
    ("predicado verdadeiro → age → unlock", "", VERDE, BRANCO),
]
for i, (a, b, f, c) in enumerate(passos):
    yy = 1.85 + i * 0.78
    caixa(s, 0.6, yy, 5.6, 0.62, fundo=f, txt=[[(a, {"negrito": True, "fonte": F_COD, "tam": 13, "cor": c})]] +
          ([[(b, {"tam": 12, "cor": c})]] if b else []), alinh=PP_ALIGN.LEFT, margem=0.2)
    if i < len(passos) - 1:
        seta(s, 3.4, yy + 0.62, 3.4, yy + 0.78, cor=CINZA_M)
seta(s, 6.2, 5.15, 6.45, 5.15, ponta=False, cor=CINZA_M)
seta(s, 6.45, 5.15, 6.45, 2.95, ponta=False, cor=CINZA_M)
seta(s, 6.45, 2.95, 6.2, 2.95, cor=CINZA_M)
codigo(s, 6.9, 1.85, 5.8, "rena.c", [(36, 38)], dest={37}, tam=13, titulo_cod="src/rena.c · a rena espera a viagem")
tabela(s, 6.9, 3.2, [1.6, 2.1, 2.1], [
    ["", "lock no mutex", "cond_wait"],
    ["Quer", "acesso exclusivo", "um evento"],
    ["Segura o mutex?", "ainda não", "não: soltou"],
    ["Dura", "µs", "ms"],
], tam=13, alt=0.42, negrito_col0=True)
texto(s, 6.9, 5.05, 5.8, 1.6, [
    [("Por que while e não if: ", {"negrito": True}),
     ("a thread pode acordar sem motivo, e um broadcast em cv_elfos acorda elfos de outros grupos, que precisam voltar a dormir.", {})],
], tam=13)
rodape(s, "src/rena.c:36-38; src/elfo.c:57-58; src/noel.c:100-107")
notas(s, "Q14: o que acontece quando uma thread espera e o que a faz continuar. Ela continua quando outra thread muda "
         "o estado E avisa (ou no timeout de Noel). O aviso não carrega informação: quem decide é o estado relido "
         "com o mutex.")

# 13 — grupo de 3 sem invasão ------------------------------------------------------
s = novo()
titulo(s, "O 3º elfo fecha o grupo na mesma seção crítica, por isso um 4º elfo nunca entra", "Sincronização")
codigo(s, 0.6, 1.85, 6.6, "elfo.c", [(27, 27), (34, 36), (41, 41), (45, 45), (47, 47), (49, 51), (53, 53)], dest={41, 47, 49, 50, 51, 53}, tam=12,
       titulo_cod="src/elfo.c · tudo entre o lock (27) e o wait (58)")
rotulo(s, 7.5, 1.85, 5.2, "6 elfos, 2 grupos: estado após cada elfo")
tabela(s, 7.5, 2.2, [0.75, 1.6, 0.85, 0.85, 1.15], [
    ["Elfo", "Ação", "g0.n", "g1.n", "aberto"],
    ["0", "entra em g0", "1", "0", "0"],
    ["1", "entra em g0", "2", "0", "0"],
    [[("2", {"negrito": True})], [("fecha g0 · signal", {"negrito": True})], "3", "0", [("1", {"negrito": True, "cor": ACENTO})]],
    ["3", "entra em g1", "3", "1", "1"],
    ["4", "entra em g1", "3", "2", "1"],
    ["5", "fecha g1 · signal", "3", "3", "1*"],
], tam=13, alt=0.42, destaque_linhas=(3,))
texto(s, 7.5, 5.2, 5.2, 0.5, "* último grupo: grupo_aberto não avança (elfo.c:50)", tam=11, cor=CINZA)
callout(s, 0.6, 5.85, 12.1, 0.85,
        "Inscrever, marcar PRONTO, contar o grupo e avançar grupo_aberto acontecem sem soltar o mutex: não existe "
        "instante em que o elfo 3 veja g0 com 3 membros e ainda aberto.", tam=14)
rodape(s, "src/elfo.c:27-58")
notas(s, "Pergunta 'como impedir que um 4º elfo entre'. Fila de grupos com id em vez de um contador global: cada "
         "elfo só entra em grupos[grupo_aberto]. O if da linha 36 é asserção defensiva: com o código atual nunca "
         "dispara. Noel só atende grupos entre proximo_grupo e grupos_fechados (noel.c:90).")

# 14 — 9ª rena e liberação -----------------------------------------------------------
s = novo()
titulo(s, "A 9ª rena acorda Noel; ao fim da preparação, um único broadcast libera as nove", "Sincronização")
ref = faixas(ev, "E06_R10_rep1")
_e = ev[(ev["execucao"] == "E06_R10_rep1") & (ev["evento"] == "decisao")]
t_dec = float(_e["t_ms"].iloc[0])
lanes = [("Renas 1–8", 2.2), ("9ª rena", 3.2), ("Noel", 4.2)]
for nome, yy in lanes:
    texto(s, 0.6, yy, 1.5, 0.5, nome, tam=14, negrito=True, ancora=MSO_ANCHOR.MIDDLE, alinh=PP_ALIGN.RIGHT)
    caixa(s, 2.2, yy + 0.24, 10.5, 0.02, fundo=LINHA)
eventos14 = [
    (2.4, 2.2, "lock · renas_prontas++ · cond_wait(cv_renas)", FUNDO2, TINTA, 3.3),
    (2.4, 3.2, "lock · ++ → 9 · t_nona · signal(cv_noel)", RENA, BRANCO, 3.3),
    (5.95, 4.2, f"acorda · renas_ok = 1 · decide: viagem", VERDE, BRANCO, 2.6),
    (8.7, 4.2, "unlock · prepara 5 ms · lock", PREP, BRANCO, 2.2),
    (11.0, 3.2, "broadcast(cv_renas)", VERDE, BRANCO, 1.7),
]
for x, yy, txt, f, c, w in eventos14:
    caixa(s, x, yy, w, 0.5, fundo=f, txt=txt, tam=12, cor=c, negrito=True)
seta(s, 5.7, 3.7, 6.0, 4.2, cor=RENA, esp=1.5)
seta(s, 11.85, 4.2, 11.85, 3.7, cor=VERDE, esp=1.5)
seta(s, 11.85, 3.2, 11.85, 2.7, cor=VERDE, esp=1.5)
texto(s, 2.4, 3.75, 3.5, 0.3, f"{br(ref['nona'], 3)} ms", tam=12, cor=RENA, negrito=True)
texto(s, 5.95, 4.75, 2.6, 0.3, f"decisão em {br(t_dec, 3)} ms", tam=12, cor=VERDE, negrito=True)
texto(s, 10.0, 4.75, 2.7, 0.3, f"fim em {br(ref['prep'][0][1], 3)} ms", tam=12, cor=VERDE, negrito=True,
      alinh=PP_ALIGN.RIGHT)
codigo(s, 0.6, 5.2, 6.3, "rena.c", [(28, 28), (30, 33)], dest={33}, tam=12)
texto(s, 7.2, 5.2, 5.5, 1.5, [
    [("Só a nona avisa: ", {"negrito": True}), ("as outras oito só se contam. Noel não precisa acordar 9 vezes.", {})],
    [("Resposta às renas: ", {"negrito": True}),
     (f"{br(1000 * (t_dec - ref['nona']), 0)} µs nesta execução (média {br(1000 * r06['resposta_renas_media_ms'], 0)} µs).", {})],
], tam=13, espaco=6)
rodape(s, "src/rena.c:23-38; src/noel.c:19-42; dados/eventos.csv (E06_R10_rep1)")
notas(s, "Sequência real da referência. Ponto a explicar: o broadcast em cv_renas acorda as 9 renas; cada uma "
         "readquire o mutex, vê viagem_concluida = 1 e sai. O signal em cv_noel basta porque só Noel espera nela.")

# 15 — decisão de Noel -------------------------------------------------------------
s = novo()
titulo(s, "Noel decide sozinho, com o mutex, e testa as renas antes da fila de elfos", "Sincronização")
fl = [
    (0.6, 1.95, 3.0, 0.62, "topo do laço (mutex travado)", FUNDO2, TINTA),
    (0.6, 2.85, 3.0, 0.62, "viagem e todos os grupos\nconcluídos? → encerra", FUNDO2, TINTA),
    (0.6, 3.75, 3.0, 0.62, "lê renas_ok e elfos_ok", FUNDO2, TINTA),
    (0.6, 4.65, 3.0, 0.62, "nada pronto → timedwait\n(cv_noel), depois volta", FUNDO2, TINTA),
    (0.6, 5.55, 3.0, 0.62, "renas_ok? → viagem\nsenão → grupo mais antigo", ACENTO, BRANCO),
]
for x, yy, w, h, t, f, c in fl:
    caixa(s, x, yy, w, h, fundo=f, txt=t, tam=13, cor=c, negrito=f == ACENTO)
for yy in (2.57, 3.47, 4.37, 5.27):
    seta(s, 2.1, yy, 2.1, yy + 0.28, cor=CINZA_M)
seta(s, 3.6, 5.86, 3.85, 5.86, ponta=False, cor=CINZA_M)
seta(s, 3.85, 5.86, 3.85, 2.26, ponta=False, cor=CINZA_M)
seta(s, 3.85, 2.26, 3.6, 2.26, cor=CINZA_M)
texto(s, 3.9, 3.8, 0.9, 0.9, "após a tarefa", tam=10, cor=CINZA)
codigo(s, 4.9, 1.85, 7.8, "noel.c", [(89, 90), (100, 101), (117, 117), (123, 131)], dest={89, 90, 117, 128}, tam=12,
       titulo_cod="src/noel.c · thread_noel")
callout(s, 4.9, 5.75, 7.8, 0.95,
        "A prioridade vem da ordem do teste (linha 117), não do mutex. O mutex só garante que Noel decide sobre um "
        "retrato consistente: ninguém muda o estado entre ler e decidir.", tam=13)
rodape(s, "src/noel.c:85-132")
notas(s, "Pergunta 'como a prioridade das renas é garantida'. Duas razões: (1) só a thread Noel decide, lendo as "
         "duas demandas com o mutex; (2) renas_ok é testado primeiro. Ponto sutil: uma rena bloqueada no lock no "
         "instante da decisão ainda não chegou para o programa; isso não é violação.")

# 16 — ajuda não é interrompida -------------------------------------------------------
s = novo()
titulo(s, "A ajuda não é interrompida: a rena que chega durante ela só é vista no topo do laço", "Sincronização")
codigo(s, 0.6, 1.85, 5.2, "noel.c", [(58, 62)], dest={60, 61, 62}, tam=12, titulo_cod="src/noel.c · ajudar_grupo")
texto(s, 6.1, 1.85, 6.6, 1.5, [
    [("Durante a ajuda, Noel solta o mutex e só dorme. ", {"negrito": True}),
     ("Ele não olha nada: elfos e renas continuam chegando e gravando o estado.", {})],
    [("O signal da 9ª rena se perde ", {"negrito": True}),
     ("(Noel não está em cv_noel), mas renas_prontas = 9 fica gravado e é relido na próxima decisão.", {})],
], tam=13, espaco=6)
R = R13
_, px = linha_tempo(s, 0.6, 4.3, 12.1, 11.5, 24.0, [
    {"rot": "Noel", "neg": True,
     "segs": [(R["dec1"], R["fim1"], AJUDA, "ajuda grupo 1"), (R["dec2"], R["fimprep"], PREP, "viagem"),
              (R["dec3"], R["fim3"], AJUDA, "ajuda grupo 2")],
     "marcas": [(R["nona"], RENA, f"9ª rena {br(R['nona'], 3)}", 1),
                (R["g2_pronto"], PRONTO, f"grupo 2 pronto {br(R['g2_pronto'], 3)}", 1),
                (R["dec2"], ACENTO, f"decisão {br(R['dec2'], 3)}: viagem", 2)]},
], alt=0.6, ticks=[12, 14, 16, 18, 20, 22, 24], rot_w=1.0)
callout(s, 0.6, 5.65, 12.1, 1.05,
        f"A 9ª rena chegou em {br(R['nona'], 3)} ms e Noel só reagiu em {br(R['dec2'], 3)} ms, depois de terminar "
        "a ajuda. Na volta, as duas demandas estavam prontas, e a viagem venceu: o grupo 2 esperou mais 5 ms.", tam=14)
rodape(s, "src/noel.c:58-62, 89-90; verificação em 2026-10-28: ./simulacao --elfos 6 --renas-ms 13 (fora da matriz)")
notas(s, "Prioridade não preemptiva. Pergunta típica: 'Noel solta o mutex na ajuda; outra thread não pode começar "
         "outra ajuda?' Não: só a thread Noel executa ajudar_grupo/preparar_viagem; o grupo em ajuda está EM_AJUDA "
         "e já saiu da fila (proximo_grupo avançou). Este cenário (--renas-ms 13) não faz parte da matriz oficial.")

# 17 — antes / junto / depois ----------------------------------------------------------
s = novo()
titulo(s, "Antes, junto ou depois: o horário das renas muda quem espera, nunca as regras", "Sincronização")
cen = [("E06_R05_rep1", "Antes (5 ms)", a05, "Elfos não esperam pelas renas"),
       ("E06_R10_rep1", "Junto (10 ms)", r06, "Grupo 1 espera a viagem inteira"),
       ("E06_R30_rep1", "Depois (30 ms)", d30, "Noel fica ocioso até as renas")]
yy = 2.15
for k, (ex, rot, t, msg) in enumerate(cen):
    f = faixas(ev, ex)
    segs = [(a, b, PREP, "viagem") for a, b, _ in f["prep"]] + [(a, b, AJUDA, f"g{int(g)}") for a, b, g in f["ajuda"]]
    marcas = [(f["nona"], RENA, "", 0)] + [(p, PRONTO, "", 0) for p in f["prontos"]]
    linha_tempo(s, 0.6, yy, 8.6, 4, 36, [{"rot": rot, "neg": k == 1, "segs": segs, "marcas": marcas}],
                alt=0.55, ticks=[5, 10, 15, 20, 25, 30, 35] if k == 2 else None, rot_w=1.7, tam=10)
    texto(s, 9.5, yy - 0.1, 3.2, 0.8, [[(msg, {"negrito": True, "cor": ACENTO if k == 1 else TINTA})],
                                      f"p95 da espera: {br(t['espera_p95_ms_media_execucoes'])} ms"], tam=13, espaco=2)
    yy += 1.25
legenda_cores(s, 2.3, 5.95, [(PREP, "viagem (5 ms)"), (AJUDA, "ajuda (3 ms)"), (RENA, "9ª rena"), (PRONTO, "grupo pronto")])
texto(s, 0.6, 6.4, 12.1, 0.5, "Grupos ficam prontos em ≈12 e ≈15 ms. Só no caso “junto” a viagem cai em cima da "
      "formação do 1º grupo, e é aí que os elfos pagam a prioridade.", tam=13, cor=CINZA)
rodape(s, "dados/eventos.csv (repetição 1 de E06_R05, E06_R10 e E06_R30); dados/tabela_esperas.csv")
notas(s, "Os três cenários da matriz, com dados reais. Liga a sincronização ao G2: o p95 sobe só no caso 'junto'. "
         "No caso 'antes', o grupo 2 fica pronto durante a ajuda 1 e é atendido logo depois, sem dormir, porque "
         "elfos_ok já era 1 no topo do laço.")

# 18 — correta × ingênua ----------------------------------------------------------------
s = novo()
titulo(s, "Com os mesmos mecanismos, só a regra de decisão separa a versão correta da ingênua", "Sincronização")
for k, (df, nome, r) in enumerate([(ev_tp, "Correta", tp), (ev_in, "Ingênua", ing)]):
    ex = df["execucao"].iloc[0]
    f = faixas(df, ex)
    segs = [(a, b, PREP, "viagem") for a, b, _ in f["prep"]] + [(a, b, AJUDA, f"g{int(g)}") for a, b, g in f["ajuda"]]
    marcas = [(f["nona"], RENA, f"9ª rena {br(f['nona'], 1)}" if k == 0 else "", 1)] + \
             [(p, PRONTO, f"g{i + 1} pronto" if k == 0 else "", 1 + i) for i, p in enumerate(f["prontos"])]
    linha_tempo(s, 0.6, 2.45 + k * 1.15, 8.3, 10, 43, [{"rot": nome, "neg": True, "segs": segs, "marcas": marcas}],
                alt=0.55, ticks=[10, 15, 20, 25, 30, 35, 40] if k == 1 else None, rot_w=1.3, tam=11)
tabela(s, 9.2, 1.95, [1.55, 1.0, 1.0], [
    ["", "Correta", "Ingênua"],
    ["1ª decisão", "renas", [("elfos", {"negrito": True, "cor": ACENTO})]],
    ["Violações", str(int(tp["violacoes_prioridade"])), [(str(int(ing["violacoes_prioridade"])), {"negrito": True, "cor": ACENTO})]],
    ["Resposta às renas", f"{br(1000 * tp['resposta_renas_ms'], 0)} µs", f"{br(ing['resposta_renas_ms'], 1)} ms"],
    ["Resultado", "PASSOU", "FALHOU"],
], tam=13, alt=0.45, negrito_col0=True)
codigo(s, 0.6, 4.95, 8.3, "noel.c", [(117, 119)], dest={117}, tam=12)
texto(s, 9.2, 4.4, 3.5, 1.4, "--teste-prioridade segura Noel até as duas demandas estarem prontas (noel.c:94-98), "
      "forçando a disputa que a matriz não produz.", tam=12, cor=CINZA)
callout(s, 0.6, 5.95, 12.1, 0.75, "Ter mutex e cond vars não garante a prioridade: a versão ingênua usa os mesmos "
        "mecanismos e viola a regra. O detector de violações funciona, já que contou 2.", tam=14)
rodape(s, "dados/teste_prioridade/ e dados/ingenua/ (eventos.csv, resumo.csv, saida.txt); src/noel.c:94-98, 117-119")
notas(s, "Comparação lado a lado com o mesmo cenário (6 elfos, renas em 30 ms). Correta: viagem primeiro e depois "
         "g1 e g2. Ingênua: g1, g2 e só então a viagem; as renas esperam ~6 ms em vez de ~20 µs.")

# 19 — critérios ----------------------------------------------------------------------
s = novo()
titulo(s, "A solução atende aos critérios clássicos; o ponto fraco é a robustez, não a sincronização", "Propriedades")
tabela(s, 0.6, 1.85, [2.9, 1.95, 7.25], [
    ["Critério", "Atende?", "Evidência"],
    ["Exclusão mútua", "sim", "todo estado mutável só com o mutex"],
    ["Progresso", "sim", "se renas_ok ou elfos_ok, Noel decide (noel.c:100-132)"],
    ["Espera limitada", "sim (finita)", "fila FIFO de grupos e no máximo 1 viagem de 5 ms"],
    ["Sem hipótese de velocidade", "sim", "toda espera é por predicado; os sleeps só simulam durações"],
    ["Sem espera ocupada", "sim", "todas as esperas são cond_wait / timedwait"],
    ["Sem deadlock", "sim", "1 mutex, nenhum lock aninhado"],
    [[("Robustez", {"negrito": True})], [("parcial", {"negrito": True, "cor": ACENTO})],
     "retornos de init não verificados (main.c:131-139); pthread_create falho sai sem joins"],
    ["Desempenho", "bom p/ N pequeno", "o mutex único serializa tudo; broadcast em cv_elfos acorda outros grupos"],
], tam=14, alt=0.5, negrito_col0=True, destaque_linhas=(7,))
texto(s, 0.6, 6.45, 12.1, 0.4, "Melhorias: verificar os retornos; uma cond var por grupo; prazo na espera da largada "
      "(main.c:208).", tam=13, cor=CINZA)
rodape(s, "src/; vault: Riscos de Sincronização (Papai Noel)")
notas(s, "Q6. Falar com honestidade: a sincronização atende; a robustez é parcial. Um mutex único não é gargalo "
         "aqui (22 threads, seções críticas de microssegundos), mas com milhares de threads valeria separar locks, "
         "com cuidado com a ordem de aquisição.")

# 20 — deadlock / starvation -----------------------------------------------------------
s = novo()
titulo(s, "Não há deadlock nem starvation na versão finita; uma versão contínua precisaria de aging", "Propriedades")
cols = [
    ("Deadlock", "threads esperando umas pelas outras para sempre",
     ["Precisa das 4 condições de Coffman juntas, incluindo espera circular.",
      "Com 1 mutex não há ciclo entre locks; quem faz cond_wait solta o mutex.",
      "Todo predicado tem quem o torne verdadeiro e avise."], "não ocorre"),
    ("Starvation", "uma thread nunca progride porque outras passam na frente",
     ["Finita: 1 viagem de 5 ms; grupos em FIFO (proximo_grupo++, noel.c:128).",
      "Contínua: renas sempre primeiro podem deixar elfos esperando para sempre.",
      "Mitigação: aging ou cota de grupos entre viagens."], "não na versão finita"),
    ("Espera ocupada e corrida", "gastar CPU testando em laço; resultado dependente da ordem",
     ["Os sleeps só simulam durações; esperas usam cond vars.",
      "Corrida descartada por inspeção: todo acesso mutável com o mutex.",
      "TSan e Helgrind não estavam instalados: não verificado por ferramenta."], "descartadas por inspeção"),
]
for i, (nome, defi, itens, veredito) in enumerate(cols):
    x = 0.6 + i * 4.1
    caixa(s, x, 1.85, 3.9, 0.5, fundo=VERDE, txt=nome, tam=16, cor=BRANCO, negrito=True)
    texto(s, x + 0.05, 2.45, 3.8, 0.7, defi, tam=13, cor=CINZA)
    texto(s, x + 0.05, 3.2, 3.8, 2.3, [f"• {t}" for t in itens], tam=13, espaco=6)
    caixa(s, x, 5.35, 3.9, 0.45, fundo=FUNDO2, txt=veredito, tam=14, negrito=True, cor=VERDE)
callout(s, 0.6, 6.0, 12.1, 0.7, "O timeout de 2000 ms é rede de segurança, não admissão de deadlock: nunca "
        f"disparou em {N_EXEC} + {EST_N} execuções.", tam=14)
rodape(s, "src/noel.c:101-106, 128; vault: Riscos de Sincronização (Papai Noel)")
notas(s, "Q7. Deadlock: as 4 condições de Coffman (exclusão mútua, posse e espera, não preempção, espera circular). "
         "Aqui falta a espera circular. Limitação do timeout: não cobre um mutex nunca liberado, porque o timedwait "
         "precisa readquirir o mutex para retornar.")

# 21 — método ---------------------------------------------------------------------------
s = novo()
titulo(s, "25 execuções em 5 configurações isolam o efeito do nº de elfos e o do horário das renas", "Resultados")
rotulo(s, 0.6, 1.85, 5, "Matriz: elfos × chegada das renas")
xs, ys = [2.0, 3.3, 4.6], [2.65, 3.45, 4.25]
for j, r in enumerate(["5 ms", "10 ms", "30 ms"]):
    texto(s, xs[j], 2.25, 1.2, 0.3, r, tam=13, cor=CINZA, alinh=PP_ALIGN.CENTER, negrito=True)
for i, e in enumerate(["3", "6", "12"]):
    texto(s, 0.6, ys[i], 1.3, 0.7, f"{e} elfos", tam=13, cor=CINZA, alinh=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE,
          negrito=True)
usadas = {(0, 1): "C1", (1, 0): "C2", (1, 1): "REF", (1, 2): "C2", (2, 1): "C1"}
for i in range(3):
    for j in range(3):
        k = usadas.get((i, j))
        caixa(s, xs[j], ys[i], 1.2, 0.7, fundo=(ACENTO if k == "REF" else VERDE_CL) if k else FUNDO2,
              txt=("referência" if k == "REF" else ("× 5" if k else "")), tam=12,
              cor=BRANCO if k == "REF" else VERDE, negrito=True)
texto(s, 0.6, 5.1, 5.3, 1.3, ["Coluna 10 ms: conjunto 1 (varia elfos → G1).", "Linha 6 elfos: conjunto 2 (varia renas → G2).",
                              "Sementes 11–55 só identificam a repetição: sem aleatoriedade; a variação vem do SO."],
      tam=13, cor=CINZA, espaco=4)
rotulo(s, 6.4, 1.85, 6, "Métricas")
tabela(s, 6.4, 2.2, [2.0, 4.3], [
    ["Métrica", "Definição"],
    ["Vazão", "pedidos concluídos / T (s)"],
    ["Espera do elfo", "início da ajuda − pedido"],
    ["p95", "posição ⌈0,95·n⌉ das esperas ordenadas"],
    ["Resposta às renas", "início da preparação − 9ª rena"],
], tam=13, alt=0.42, negrito_col0=True)
rotulo(s, 6.4, 4.5, 6, "Testes além da matriz")
texto(s, 6.4, 4.85, 6.3, 1.8, [
    "• --teste-prioridade: força a disputa → PASSOU",
    f"• --ingenua: mesmo cenário → {int(ing['violacoes_prioridade'])} violações detectadas",
    f"• estresse: {EST_OK}/{EST_N} ok · timeout de 5 ms: status timeout, sem travar",
], tam=13, espaco=4)
rodape(s, "scripts/run_experimentos.sh; src/metricas.c:17-27, 43-71; dados/")
notas(s, "Q8 começa aqui. Para cada gráfico: o que ficou fixo, o que variou, quantas execuções, a tendência e que "
         "parte do código a explica. Atenção: com n = 3 ou 6 amostras, o p95 é o máximo.")

# 22 — G1 --------------------------------------------------------------------------------
s = novo()
titulo(s, f"Mais elfos diluem o custo fixo inicial: a vazão sobe de {br(r03['vazao_media_pedidos_s'], 0)} para "
          f"{br(r12['vazao_media_pedidos_s'], 0)} pedidos/s", "Resultados")
imagem(s, G / "G1_vazao_elfos.png", 0.6, 1.8, h=4.95)
texto(s, 8.4, 1.85, 4.3, 3.8, [
    [("Fixo: ", {"negrito": True}), ("renas em 10 ms; 5 execuções por ponto.", {})],
    [("Variou: ", {"negrito": True}), ("3, 6 e 12 elfos (eixo X).", {})],
    [("Eixo Y: ", {"negrito": True}), ("pedidos de elfos atendidos por segundo.", {})],
    [("Por quê: ", {"negrito": True}),
     ("há um custo fixo de ≈15 ms (10 ms até o 1º elfo + 5 ms de viagem). Depois, a fila FIFO atende um grupo "
      "atrás do outro: T ≈ 15 + 3·(N/3) ms.", {})],
], tam=13, espaco=8)
callout(s, 8.4, 5.7, 4.3, 1.05, "Previsto 167 / 286 / 444; medido ≈2% abaixo (latência dos sleeps). "
        "A faixa mín–máx é menor que o marcador.", tam=13)
rodape(s, "graficos/G1_vazao_elfos.png; dados/tabela_esperas.csv")
notas(s, f"G1. Valores: {br(r03['vazao_media_pedidos_s'], 1)}, {br(r06['vazao_media_pedidos_s'], 1)}, "
         f"{br(r12['vazao_media_pedidos_s'], 1)} pedidos/s. A previsão 15 + 3·N/3 é uma conta nossa sobre o cenário, "
         "não um dado medido. Limite teórico: 3 elfos a cada 3 ms = 1000 pedidos/s.")

# 23 — G2 --------------------------------------------------------------------------------
s = novo()
titulo(s, f"Só com as renas junto os elfos pagam a prioridade: p95 de {br(a05['espera_p95_ms_media_execucoes'], 1)} "
          f"para {br(r06['espera_p95_ms_media_execucoes'], 1)} ms", "Resultados")
imagem(s, G / "G2_p95_renas.png", 0.6, 1.75, w=8.6)
texto(s, 9.5, 1.85, 3.2, 3.9, [
    [("Esquerda: ", {"negrito": True}), ("p95 da espera dos elfos por cenário (categorias, sem linha).", {})],
    [("Direita: ", {"negrito": True}), ("resposta às renas, em µs.", {})],
    [("Por quê: ", {"negrito": True}),
     ("só no caso “junto” a viagem (≈10–15 ms) cai em cima da formação do 1º grupo (≈12 ms). Nos outros, Noel "
      "está livre quando cada grupo fecha.", {})],
], tam=13, espaco=8)
callout(s, 0.6, 5.65, 12.1, 1.05, f"As renas são atendidas em ≈{br(1000 * r06['resposta_renas_media_ms'], 0)}–"
        f"{br(1000 * d30['resposta_renas_media_ms'], 0)} µs em qualquer cenário: o preço da prioridade é pago só pelos "
        "elfos do 1º grupo, cerca de 3 ms a mais. A diferença é muito maior que a variação entre repetições.", tam=14)
rodape(s, "graficos/G2_p95_renas.png; dados/tabela_esperas.csv")
notas(s, "G2. A resposta às renas (~20 µs) é o custo de acordar Noel: o signal em rena.c:33 mais a reaquisição do "
         "mutex. As barras de erro do painel direito parecem grandes, mas estão em microssegundos.")

# 24 — G3 --------------------------------------------------------------------------------
s = novo()
titulo(s, f"Os dois grupos esperam o mesmo ({br(G3[1])} × {br(G3[2])} ms): o atraso da viagem se propaga sem crescer",
       "Resultados")
imagem(s, G / "G3_p95_grupo.png", 0.6, 1.8, h=4.95)
texto(s, 8.4, 1.85, 4.3, 3.8, [
    [("Fixo: ", {"negrito": True}), ("referência (6 elfos, renas junto).", {})],
    [("Variou: ", {"negrito": True}), ("o grupo atendido; 5 execuções × 3 elfos = 15 amostras por grupo.", {})],
    [("Por quê: ", {"negrito": True}),
     ("o grupo 1 espera o fim da viagem; o grupo 2 espera o fim da ajuda ao grupo 1. Como os grupos chegam a cada "
      "3 ms e cada ajuda dura 3 ms, o atraso passa adiante sem acumular.", {})],
], tam=13, espaco=8)
callout(s, 8.4, 5.6, 4.3, 1.15, "A expectativa intuitiva (“o 2º grupo espera mais”) não se confirmou, e isso faz "
        "parte da análise.", tam=13)
rodape(s, "graficos/G3_p95_grupo.png; dados/espera_elfos.csv")
notas(s, "G3. Com 3 amostras por execução, o p95 é o máximo do grupo; por isso o gráfico mostra os pontos "
         "individuais. Dizer que a hipótese não se confirmou mostra que interpretamos o dado em vez de forçá-lo.")

# 25 — G4 --------------------------------------------------------------------------------
s = novo()
titulo(s, "Na linha do tempo, os grupos ficam prontos durante a viagem e esperam sem interrompê-la", "Resultados")
imagem(s, G / "G4_linha_tempo.png", 0.6, 1.7, w=9.3)
texto(s, 10.1, 1.85, 2.6, 4.0, [
    [("Linha Noel: ", {"negrito": True}), ("estado ao longo do tempo (cinza, laranja, azul).", {})],
    [("Linhas dos grupos: ", {"negrito": True}), ("hachura = pronto e esperando; azul claro = em ajuda.", {})],
    [("Marcas: ", {"negrito": True}), ("9ª rena (roxo) e grupos prontos (verde tracejado).", {})],
], tam=13, espaco=8)
callout(s, 0.6, 5.75, 12.1, 0.95, f"Grupos prontos em {br(ref['prontos'][0])} e {br(ref['prontos'][1])} ms, "
        f"durante a viagem ({br(ref['prep'][0][0])}–{br(ref['prep'][0][1])} ms): entram na fila e a preparação "
        "segue até o fim. Depois, Noel fica ocupado sem pausa.", tam=14)
rodape(s, "graficos/G4_linha_tempo.png; dados/eventos.csv (E06_R10_rep1)")
notas(s, "G4 é a prova visual do slide 16 na rodada oficial: nenhuma faixa laranja é cortada por uma azul. "
         "É a mesma execução do slide 14.")

# 26 — observado × garantido ---------------------------------------------------------------
s = novo()
titulo(s, "Os testes mostram que as regras valeram nas execuções; a garantia vem da análise do código", "Resultados")
cols = [
    ("Observado nos dados", VERDE, [
        f"{N_OK}/{N_EXEC} execuções ok, {VIOL} violações",
        "teste de prioridade PASSOU",
        f"ingênua: {int(ing['violacoes_prioridade'])} violações → o detector funciona",
        f"estresse {EST_OK}/{EST_N} ok; timeout nunca disparou"]),
    ("Garantido por construção", VERDE, [
        "grupo de 3 sem invasão (mesma seção crítica)",
        "viagem só com 9 renas, uma vez",
        "renas primeiro sobre o estado lido",
        "tarefa iniciada não é interrompida"]),
    ("Limites da evidência", ACENTO, [
        "na matriz, 0 decisões com disputa: a prioridade só foi exercitada pelos testes dedicados",
        "algumas checagens nunca disparam por construção",
        "TSan e Helgrind não foram executados",
        "quem “chega primeiro” num empate depende do escalonador"]),
]
for i, (cab, cor, itens) in enumerate(cols):
    x = 0.6 + i * 4.1
    caixa(s, x, 1.85, 3.9, 0.5, fundo=cor, txt=cab, tam=16, cor=BRANCO, negrito=True)
    texto(s, x + 0.05, 2.5, 3.8, 3.2, [f"• {t}" for t in itens], tam=14, espaco=8)
callout(s, 0.6, 5.85, 12.1, 0.85, "Testes mostram a presença de erros, não a ausência. Por isso a defesa da "
        "solução se apoia no raciocínio sobre o código, e os dados servem de evidência.", tam=14)
rodape(s, "dados/log_execucao.txt (com disputa: 0); vault: Riscos de Sincronização; Fluxo do Código, seção 6")
notas(s, "Pergunta 'os resultados provam que a solução está correta?'. Não provam; mostram que as regras valeram "
         "nas execuções observadas. Ver slide de apoio sobre o empate natural (--renas-ms 12).")

# 27 — conclusão ------------------------------------------------------------------------------
s = novo()
titulo(s, "Um estado consistente e uma regra de decisão resolvem o problema; robustez e justiça contínua são os próximos passos")
cols = [
    ("Como funciona", ["1 mutex protege toda a Oficina", "4 cond vars: largada, Noel, elfos, renas",
                       "fila de grupos com id; o 3º elfo fecha o grupo", "só Noel decide: renas_ok primeiro"]),
    ("O que garante", ["grupos de exatamente 3, sem invasão", "viagem só com as 9 renas",
                       "prioridade às renas sem interromper a ajuda", "sem deadlock, sem espera ocupada"]),
    ("Limitações e melhorias", ["verificar retornos de pthread_*_init", "uma cond var por grupo",
                                "prazo na espera da largada", "aging numa versão contínua; rodar TSan"]),
]
for i, (cab, itens) in enumerate(cols):
    x = 0.6 + i * 4.1
    caixa(s, x, 1.85, 3.9, 0.5, fundo=ACENTO if i == 2 else VERDE, txt=cab, tam=16, cor=BRANCO, negrito=True)
    texto(s, x + 0.05, 2.5, 3.8, 3.0, [f"• {t}" for t in itens], tam=15, espaco=10)
caixa(s, 0.6, 5.55, 12.1, 1.15, fundo=FUNDO2,
      txt=[[("Mensagem final: ", {"negrito": True, "cor": VERDE}),
            ("sincronizar não é só travar: é proteger o estado, esperar pelo predicado certo e decidir sobre um "
             "retrato consistente. A prioridade vem da regra; os mecanismos garantem que ela seja aplicada.", {})]],
      tam=16, alinh=PP_ALIGN.LEFT, margem=0.25)
rodape(s, "síntese dos slides 2–26")
notas(s, "Fechar em 30 s e abrir para perguntas. Os slides de apoio a seguir servem para respostas rápidas.")

# A1 — mapa de perguntas ----------------------------------------------------------------------
s = novo()
titulo(s, "Apoio · Cada pergunta do professor tem um slide de resposta", "Apoio")
tabela(s, 0.6, 1.8, [8.3, 3.8], [
    ["Pergunta", "Slides"],
    ["1. O que é um problema de sincronização?", "3, 4"],
    ["2. Como se relaciona com o Sistema Operacional?", "3, 7"],
    ["3. O que são mutex e semáforo?", "5, 29"],
    ["4. if + contador × mutex", "4, 13"],
    ["5. Relação com a convivência de threads e processos", "7"],
    ["6. Princípios de uma boa solução; atendemos? poderia ser melhor?", "19, 27"],
    ["7. Starvation e deadlock aparecem?", "20"],
    ["8. Como apresentar os gráficos e interpretar os resultados?", "21–26"],
    ["Como o mutex começa e qual linha o inicializa?", "11"],
    ["Usa mutex, semáforo ou ambos? Valor inicial de cada mecanismo?", "10, 11, 29"],
    ["O que acontece quando uma thread espera e o que a faz continuar?", "12, 30"],
    ["Como a prioridade das renas é garantida?", "15, 16, 18"],
    ["Como impedir que um 4º elfo entre num grupo já iniciado?", "13"],
], tam=13, alt=0.36)
rodape(s, "vault: Perguntas da Apresentação (Papai Noel)")
notas(s, "Use durante as perguntas para pular direto ao slide certo.")

# A2 — semáforos (solução clássica) ------------------------------------------------------------
s = novo()
titulo(s, "Apoio · Mutex começa sem valor; semáforo começa com um número que muda o comportamento", "Apoio")
codigo(s, 0.6, 1.85, 5.6, "main.c", [(131, 131)], dest={131}, tam=13, titulo_cod="nosso código: mutex")
caixa(s, 0.6, 2.85, 5.6, 0.75, fundo=COD_BG,
      txt=[[("sem_init(&s, 0, valor);  ", {"fonte": F_COD, "cor": COD_TX, "tam": 13}),
            ("// não existe no projeto", {"fonte": F_COD, "cor": COD_NUM, "tam": 12})]], alinh=PP_ALIGN.LEFT, margem=0.15)
texto(s, 0.6, 3.75, 5.6, 2.2, [
    "• pshared = 0: entre threads do processo; ≠ 0: entre processos.",
    "• valor 1 → tranca; 0 → a 1ª espera bloqueia; N → N threads juntas.",
    "• No nosso código, a “memória” fica nos predicados zerados pelo memset: com renas_prontas começando em 1, "
    "a viagem sairia com 8 renas.",
], tam=13, espaco=6)
rotulo(s, 6.6, 1.85, 6, "Solução clássica de Downey (não é o nosso código)")
tabela(s, 6.6, 2.2, [1.7, 0.75, 3.65], [
    ["Semáforo", "Inicial", "Se começasse diferente"],
    ["mutex", "1", "0: ninguém entra (trava); 2: perde a exclusão"],
    ["santaSem", "0", "1: Noel acorda sem demanda"],
    ["reindeerSem", "0", "> 0: renas saem antes da preparação"],
    ["elfTex", "1", "0: nenhum elfo entra; 2: dois grupos em paralelo"],
], tam=13, alt=0.55, negrito_col0=True)
rodape(s, "A. B. Downey, The Little Book of Semaphores, “The Santa Claus problem”; src/main.c:109, 131")
notas(s, "Q12 e Q13. Deixar claro que a tabela da direita é da solução clássica, para responder 'e se o valor "
         "inicial fosse diferente'. Semáforo com 1 funciona como exclusão mútua, mas não tem dono.")

# A3 — mapa de linhas -----------------------------------------------------------------------------
s = novo()
titulo(s, "Apoio · Linhas que travam, esperam e avisam, arquivo por arquivo", "Apoio")
tabela(s, 0.6, 1.8, [2.4, 2.8, 3.3, 3.6], [
    ["Arquivo", "lock / unlock", "espera (while + wait)", "avisa"],
    ["main.c", "207 / 213", "29-30 (threads); 208-209 (main)", "broadcast 28, 212"],
    ["noel.c", "81 / 145; 31 ↔ 33; 60 ↔ 62", "100-101 timedwait cv_noel", "broadcast 41, 73, 143-144"],
    ["elfo.c", "21 / 23; 27 / 29 · 38 · 59", "57-58 cv_elfos", "signal 53"],
    ["rena.c", "17 / 19; 23 / 25 · 38", "36-37 cv_renas", "signal 33"],
], tam=13, alt=0.55, negrito_col0=True)
texto(s, 0.6, 4.75, 12.1, 1.9, [
    [("Regra seguida em todos: ", {"negrito": True}),
     ("alterar o estado com o mutex, avisar depois de alterar e testar o predicado antes de dormir.", {})],
    [("No par 31↔33 / 60↔62: ", {"negrito": True}),
     ("Noel solta o mutex só para dormir a duração da tarefa e o retoma antes de concluir.", {})],
], tam=14, espaco=8)
rodape(s, "src/main.c, noel.c, elfo.c, rena.c (conferido em 2026-10-28)")
notas(s, "Para 'quais linhas adquirem, liberam ou sinalizam'. O par 60/62 é o que deixa elfos e renas trabalharem "
         "durante a ajuda.")

# A4 — empate natural -----------------------------------------------------------------------------
s = novo()
titulo(s, "Apoio · Num empate real de chegada, quem pega o mutex primeiro define o que Noel enxerga", "Apoio")
total = EMPATE["n"]
for i, (rot, n, cor) in enumerate([("1ª decisão = viagem", EMPATE["viagem"], VERDE),
                                   ("1ª decisão = grupo 1", EMPATE["grupo"], CINZA_M)]):
    yy = 2.2 + i * 0.9
    texto(s, 0.6, yy, 2.6, 0.6, rot, tam=14, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
    caixa(s, 3.3, yy + 0.08, 5.5 * n / total, 0.45, fundo=cor)
    texto(s, 3.4 + 5.5 * n / total, yy, 1.6, 0.6, f"{n}/{total}", tam=14, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
texto(s, 0.6, 4.1, 12.1, 1.6, [
    "• Cenário: --renas-ms 12, quando a 9ª rena e o 3º elfo acordam no mesmo instante. 30 execuções, todas ok e "
    "com 0 violações.",
    "• Quando o grupo vem primeiro, a 9ª rena ainda estava na fila do mutex: para o programa, ela não tinha chegado. "
    "Na decisão seguinte, a viagem vence.",
    "• A proporção varia de uma rodada para outra (a nota do vault registra 8/30 viagem). Isso é do escalonador, não da solução.",
], tam=14, espaco=8)
callout(s, 0.6, 5.95, 12.1, 0.75, "A prioridade vale sobre o estado registrado com o mutex, não sobre o instante "
        "físico em que a thread acordou.", tam=14)
rodape(s, "verificação em 2026-10-28: 30 × ./simulacao --elfos 6 --renas-ms 12 (fora da matriz oficial)")
notas(s, "Use se o professor perguntar 'e se chegarem exatamente juntos?'. A resposta honesta: depende de quem "
         "adquire o mutex primeiro; em qualquer caso, nenhuma violação.")

prs.save(SAIDA)
print(f"{SAIDA} · {NUM[0]} slides")
